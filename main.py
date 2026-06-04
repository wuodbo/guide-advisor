import customtkinter as ctk
from tkinter import messagebox, filedialog
import pymysql
import hashlib
import random
from datetime import datetime, timedelta
from tkinter import ttk
import os
import sys
from config import DB_CONFIG


# Настройка поддержки русских шрифтов для pdf
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch, mm
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping
    
    # Регистрируем системные шрифты с кириллицей
    FONT_REGISTERED = False
    FONT_NAME = 'Helvetica'
    FONT_BOLD = 'Helvetica-Bold'
    
    # Windows шрифты
    windows_fonts = [
        ('C:/Windows/Fonts/arial.ttf', 'Arial', 'Arial'),
        ('C:/Windows/Fonts/arialbd.ttf', 'Arial-Bold', 'Arial-Bold'),
        ('C:/Windows/Fonts/times.ttf', 'Times-Roman', 'Times-Roman'),
        ('C:/Windows/Fonts/calibri.ttf', 'Calibri', 'Calibri'),
    ]
    
    # macOS шрифты
    mac_fonts = [
        ('/System/Library/Fonts/Helvetica.ttc', 'Helvetica', 'Helvetica'),
        ('/System/Library/Fonts/Arial.ttf', 'Arial', 'Arial'),
        ('/System/Library/Fonts/Supplemental/Arial.ttf', 'Arial', 'Arial'),
    ]
    
    # Linux шрифты
    linux_fonts = [
        ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 'DejaVuSans', 'DejaVuSans'),
        ('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf', 'LiberationSans', 'LiberationSans'),
        ('/usr/share/fonts/truetype/ubuntu/Ubuntu-Regular.ttf', 'Ubuntu', 'Ubuntu'),
    ]
    
    all_fonts = windows_fonts + mac_fonts + linux_fonts
    
    for font_path, font_name, font_bold in all_fonts:
        if os.path.exists(font_path):
            try:
                pdfmetrics.registerFont(TTFont(font_name, font_path))
                FONT_REGISTERED = True
                FONT_NAME = font_name
                FONT_BOLD = font_name
                
                break
            except:
                continue
    
    if not FONT_REGISTERED:
        pass

    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    FONT_NAME = 'Helvetica'
    FONT_BOLD = 'Helvetica-Bold'
    

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import BarChart, Reference, PieChart
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    

# Настройка темы в стиле Wireframe
ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

BG_MAIN = "#E5E5E5"
BG_CARD = "#D4CFC9"
BG_WHITE = "#F5F5F7"
COLOR_ACCENT = "#A39E93"
COLOR_INPUT = "#Ffffff"  # Белый цвет для полей ввода


class DB:
    @staticmethod
    def query(sql, args=None, fetch='all'):
        try:
            conn = pymysql.connect(**DB_CONFIG)
            with conn.cursor() as cur:
                cur.execute(sql, args or ())
                if fetch == 'all':
                    result = cur.fetchall()
                elif fetch == 'one':
                    result = cur.fetchone()
                else:
                    result = None
                conn.commit()
                return result
        except pymysql.Error as e:
            messagebox.showerror("Ошибка БД", str(e))
            return [] if fetch == 'all' else None

    @staticmethod
    def hash(pw):
        return hashlib.sha256(pw.encode()).hexdigest()


class PhotoPlaceholder(ctk.CTkFrame):
    def __init__(self, parent, width=150, height=150, text="Фото"):
        super().__init__(parent, width=width, height=height, fg_color=COLOR_ACCENT, corner_radius=8)
        self.pack_propagate(False)
        ctk.CTkLabel(self, text="📷", font=("Arial", 32)).place(relx=0.5, rely=0.4, anchor="center")
        ctk.CTkLabel(self, text=text, font=("Arial", 11), text_color="#333").place(relx=0.5, rely=0.7, anchor="center")


class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Guide-Advisor")
        self.geometry("1200x800")
        self.configure(fg_color=BG_MAIN)
        
        self.current_user = None
        self.current_frame = None
        
        self._init_database()
        self._seed_data()
        self._init_additional_tables()
        self.show_login()

    def _init_database(self):
        """Создание таблиц если их нет"""
        try:
            conn = pymysql.connect(
                host=DB_CONFIG['host'],
                user=DB_CONFIG['user'],
                password=DB_CONFIG['password'],
                port=DB_CONFIG['port'],
                charset='utf8mb4'
            )
            with conn.cursor() as cur:
                cur.execute("CREATE DATABASE IF NOT EXISTS guideadvisor")
                cur.execute("USE guideadvisor")
                
                # Таблица users
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS users (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        username VARCHAR(50) UNIQUE NOT NULL,
                        email VARCHAR(100) UNIQUE NOT NULL,
                        password_hash VARCHAR(255) NOT NULL,
                        first_name VARCHAR(50),
                        last_name VARCHAR(50),
                        role ENUM('ADM','GUI','TUR') DEFAULT 'TUR',
                        status ENUM('ACT','BLK','PEN') DEFAULT 'ACT',
                        registered_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                    )
                """)
                
                # Таблица guides
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS guides (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        user_id INT UNIQUE NOT NULL,
                        specialization VARCHAR(100),
                        experience_years INT DEFAULT 0,
                        description TEXT,
                        rating DECIMAL(2,1) DEFAULT 0,
                        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                    )
                """)
                
                # Таблица excursions
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS excursions (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        guide_id INT NOT NULL,
                        title VARCHAR(200) NOT NULL,
                        description TEXT,
                        category VARCHAR(50),
                        max_participants INT DEFAULT 10,
                        duration_hours INT DEFAULT 2,
                        status ENUM('PUB','HID','ARC') DEFAULT 'PUB',
                        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        FOREIGN KEY (guide_id) REFERENCES users(id)
                    )
                """)
                
                # Таблица schedules
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS schedules (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        excursion_id INT NOT NULL,
                        start_datetime DATETIME NOT NULL,
                        FOREIGN KEY (excursion_id) REFERENCES excursions(id) ON DELETE CASCADE
                    )
                """)
                
                # Таблица bookings
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS bookings (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        tourist_id INT NOT NULL,
                        schedule_id INT NOT NULL,
                        participants_count INT DEFAULT 1,
                        status ENUM('PEND','CONF','CANC','COM') DEFAULT 'PEND',
                        booked_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        FOREIGN KEY (tourist_id) REFERENCES users(id),
                        FOREIGN KEY (schedule_id) REFERENCES schedules(id)
                    )
                """)
                
                conn.commit()
            conn.close()
            print("✅ Успешное подключение к базе данных")
        except Exception as e:
            print(f"Ошибка подключения к базе данных: {e}")

    def _init_additional_tables(self):
        """Создание дополнительных таблиц"""
        try:
            conn = pymysql.connect(**DB_CONFIG)
            with conn.cursor() as cur:
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS notifications (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        user_id INT NOT NULL,
                        title VARCHAR(200) NOT NULL,
                        message TEXT NOT NULL,
                        type ENUM('INFO','BOOKING','REMINDER','PROMO') DEFAULT 'INFO',
                        is_read BOOLEAN DEFAULT FALSE,
                        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                    )
                """)
                
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS reviews (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        excursion_id INT NOT NULL,
                        user_id INT NOT NULL,
                        rating INT DEFAULT 5,
                        comment TEXT,
                        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        FOREIGN KEY (excursion_id) REFERENCES excursions(id) ON DELETE CASCADE,
                        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                    )
                """)
                
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS user_preferences (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        user_id INT NOT NULL UNIQUE,
                        favorite_categories TEXT,
                        notification_enabled BOOLEAN DEFAULT TRUE,
                        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                    )
                """)
                
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS user_activity (
                        id INT PRIMARY KEY AUTO_INCREMENT,
                        user_id INT NOT NULL,
                        excursion_id INT,
                        action_type ENUM('VIEW','BOOK','REVIEW','CANCEL') DEFAULT 'VIEW',
                        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
                    )
                """)
                
                # Добавляем колонку photo_path если её нет
                try:
                    cur.execute("ALTER TABLE guides ADD COLUMN photo_path VARCHAR(500) DEFAULT NULL")
                    conn.commit()
                except Exception:
                    pass  # Колонка уже существует
                try:
                    cur.execute("ALTER TABLE users ADD COLUMN photo_path VARCHAR(500) DEFAULT NULL")
                    conn.commit()
                except Exception:
                    pass  # Колонка уже существует
                try:
                    cur.execute("ALTER TABLE excursions ADD COLUMN photo_path VARCHAR(500) DEFAULT NULL")
                    conn.commit()
                except Exception:
                    pass  # Колонка уже существует

                conn.commit()
            conn.close()
            
        except Exception as e:
            pass

    def _seed_data(self):
        """Начальное заполнение БД тестовыми данными"""
        users_count = DB.query("SELECT COUNT(*) as c FROM users", fetch='one')
        if users_count and users_count['c'] > 0:
            return
        
        
        
        # Админ
        DB.query(
            "INSERT INTO users (username, email, password_hash, first_name, last_name, role, status)"
            " VALUES (%s,%s,%s,%s,%s,'ADM','ACT')",
            ('admin', 'admin@guide.ru', DB.hash('admin123'), 'Admin', 'System'), fetch=None
        )
        
        
        # Гиды
        guides_data = [
            ('guide_anna', 'anna@guide.ru', 'Анна', 'Сидорова', 'Исторические экскурсии', 5, 'Опытный гид по Москве, специализируется на истории Кремля и Красной площади.'),
            ('guide_mikhail', 'mikhail@guide.ru', 'Михаил', 'Ковалев', 'Архитектурные экскурсии', 3, 'Архитектор по образованию. Знает все тайны московских зданий.'),
            ('guide_elena', 'elena@guide.ru', 'Елена', 'Волкова', 'Гастрономические экскурсии', 4, 'Гастроном-эксперт. Проводит дегустации и кулинарные мастер-классы.'),
        ]
        
        for username, email, first, last, spec, exp, desc in guides_data:
            DB.query(
                "INSERT INTO users (username, email, password_hash, first_name, last_name, role, status)"
                " VALUES (%s,%s,%s,%s,%s,'GUI','ACT')",
                (username, email, DB.hash('guide123'), first, last), fetch=None
            )
            user = DB.query("SELECT id FROM users WHERE email=%s", (email,), fetch='one')
            DB.query(
                "INSERT INTO guides (user_id, specialization, experience_years, description, rating)"
                " VALUES (%s,%s,%s,%s,%s)",
                (user['id'], spec, exp, desc, round(random.uniform(4.0, 5.0), 1)), fetch=None
            )
        
        
        # Турист
        DB.query(
            "INSERT INTO users (username, email, password_hash, first_name, last_name, role, status)"
            " VALUES (%s,%s,%s,%s,%s,'TUR','ACT')",
            ('tourist', 'tourist@test.ru', DB.hash('tourist123'), 'Петр', 'Петров'), fetch=None
        )
        
        
        # Получаем всех гидов для привязки экскурсий
        guides = DB.query("SELECT user_id FROM guides", fetch='all')
        
        excursions_data = [
            ('Тайны Московского Кремля', 'Увлекательная экскурсия по историческому центру Москвы. Вы узнаете о строительстве Кремля, его тайнах и легендах.', 'Историческая', 15, 3),
            ('Шедевры московской архитектуры', 'Экскурсия по самым красивым зданиям Москвы: от средневековых палат до сталинских высоток.', 'Архитектурная', 12, 3),
            ('Гастрономический тур по Москве', 'Дегустация московских угощений. Посетим лучшие рестораны и рынки.', 'Гастрономическая', 10, 4),
            ('Парк Царицыно: природа и архитектура', 'Прогулка по парку Царицыно с уникальным ансамблем архитектуры.', 'Природная', 15, 3),
            ('Ночная Москва', 'Романтическая ночная экскурсия по огням столицы.', 'Обзорная', 20, 2),
            ('Бункер Сталина', 'Секретный бункер времен Второй мировой войны.', 'Историческая', 10, 4),
        ]
        
        for i, (title, desc, cat, maxp, dur) in enumerate(excursions_data):
            guide_id = guides[i % len(guides)]['user_id']
            DB.query(
                "INSERT INTO excursions (guide_id, title, description, category, max_participants, duration_hours, status)"
                " VALUES (%s,%s,%s,%s,%s,%s,'PUB')",
                (guide_id, title, desc, cat, maxp, dur), fetch=None
            )
        
        
        # Добавляем расписания
        excursions = DB.query("SELECT id FROM excursions", fetch='all')
        for exc in excursions:
            for day in range(1, 4):
                start_time = datetime.now() + timedelta(days=random.randint(1, 14))
                DB.query(
                    "INSERT INTO schedules (excursion_id, start_datetime) VALUES (%s, %s)",
                    (exc['id'], start_time), fetch=None
                )
        
        
        # Добавляем тестовые бронирования
        tourist = DB.query("SELECT id FROM users WHERE role='TUR'", fetch='one')
        schedules = DB.query("SELECT id FROM schedules LIMIT 3", fetch='all')
        for s in schedules:
            DB.query(
                "INSERT INTO bookings (tourist_id, schedule_id, participants_count, status)"
                " VALUES (%s,%s,%s,'CONF')",
                (tourist['id'], s['id'], random.randint(1, 4)), fetch=None
            )
        
        
        # Добавляем тестовые уведомления
        DB.query(
            "INSERT INTO notifications (user_id, title, message, type, is_read) VALUES (%s, %s, %s, %s, FALSE)",
            (tourist['id'], 'Добро пожаловать!', 'Рады приветствовать вас в Guide-Advisor!', 'INFO'), fetch=None
        )
        
        # Добавляем тестовые отзывы
        for exc in excursions[:3]:
            DB.query(
                "INSERT INTO reviews (excursion_id, user_id, rating, comment) VALUES (%s, %s, %s, %s)",
                (exc['id'], tourist['id'], random.randint(4, 5), 'Отличная экскурсия! Очень понравилось!'), fetch=None
            )
        
        

    def add_notification(self, user_id, title, message, type='INFO'):
        DB.query(
            "INSERT INTO notifications (user_id, title, message, type, is_read) VALUES (%s, %s, %s, %s, FALSE)",
            (user_id, title, message, type), fetch=None
        )

    def log_activity(self, user_id, excursion_id, action_type):
        DB.query(
            "INSERT INTO user_activity (user_id, excursion_id, action_type) VALUES (%s, %s, %s)",
            (user_id, excursion_id, action_type), fetch=None
        )

    def show_page(self, frame_class, **kwargs):
        if self.current_frame:
            self.current_frame.destroy()
        self.current_frame = frame_class(self, **kwargs)
        self.current_frame.pack(fill="both", expand=True, padx=20, pady=20)
        self.update_idletasks()
        self.current_frame.lift()

    def show_login(self):
        self.show_page(LoginPage)
    
    def show_register(self):
        self.show_page(RegisterPage)
    
    def show_forgot(self):
        self.show_page(ForgotPasswordPage)
    
    def show_dashboard(self):
        role = self.current_user['role']
        if role == 'ADM':
            self.show_page(AdminDashboardPage)
        elif role == 'GUI':
            self.show_page(GuideDashboardPage)
        else:
            self.show_page(TouristDashboardPage)
    
    def show_guides_catalog(self):
        self.show_page(GuidesCatalogPage)
    
    def show_excursions_catalog(self):
        self.show_page(ExcursionsCatalogPage)
    
    def show_guide_profile(self, guide_id):
        self.show_page(GuideProfilePage, guide_id=guide_id)
    
    def show_tourist_profile(self):
        self.show_page(TouristProfilePage)
    
    def show_manage_data(self):
        self.show_page(ManageDataPage)
    
    def show_add_excursion(self, excursion_id=None):
        self.show_page(AddExcursionPage, excursion_id=excursion_id)
    
    def show_add_guide(self, guide_id=None):
        self.show_page(AddGuidePage, guide_id=guide_id)
    
    def show_edit_excursion(self, excursion_id):
        self.show_page(EditExcursionPage, excursion_id=excursion_id)
    
    def show_edit_guide(self, guide_id):
        self.show_page(EditGuidePage, guide_id=guide_id)
    
    def show_search_catalog(self):
        self.show_page(SearchCatalogPage)
    
    def show_reviews(self, excursion_id=None):
        self.show_page(ReviewsPage, excursion_id=excursion_id)
    
    def show_guide_reviews(self, all_reviews=False):
        self.show_page(GuideReviewsPage, all_reviews=all_reviews)
    
    def show_schedule(self):
        self.show_page(SchedulePage)
    
    def show_personal_recommendations(self):
        self.show_page(PersonalRecommendationsPage)
    
    def show_notifications(self):
        self.show_page(NotificationsPage)
    
    def show_reports(self):
        self.show_page(ReportsPage)


# Логин
class LoginPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_CARD, corner_radius=12)
        self.master = master
        
        form = ctk.CTkFrame(self, fg_color="white", width=400, height=450, corner_radius=12)
        form.place(relx=0.5, rely=0.5, anchor="center")
        form.pack_propagate(False)
        
        ctk.CTkLabel(form, text="Авторизация", font=("Arial", 24, "bold")).pack(pady=(35, 25))
        
        self.email = ctk.CTkEntry(form, placeholder_text="Email", width=320, height=45, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.email.pack(pady=10)
        
        self.pwd = ctk.CTkEntry(form, placeholder_text="Пароль", show="*", width=320, height=45, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.pwd.pack(pady=10)
        
        self.err_lbl = ctk.CTkLabel(form, text="", text_color="red", font=("Arial", 12))
        self.err_lbl.pack()
        
        ctk.CTkButton(form, text="Войти", fg_color=BG_CARD, text_color="black", hover_color=COLOR_ACCENT, 
                      width=320, height=45, corner_radius=8, command=self.handle_login).pack(pady=20)
        
        sub_menu = ctk.CTkFrame(form, fg_color="transparent")
        sub_menu.pack(pady=10)
        ctk.CTkButton(sub_menu, text="Регистрация", fg_color="transparent", text_color="black", width=120,
                      command=self.master.show_register).pack(side="left", padx=10)
        ctk.CTkButton(sub_menu, text="Забыли пароль?", fg_color="transparent", text_color="black", width=130,
                      command=self.master.show_forgot).pack(side="left", padx=10)

    def handle_login(self):
        email = self.email.get()
        pwd = self.pwd.get()
        
        if not email or not pwd:
            self.err_lbl.configure(text="Введите email и пароль")
            return
        
        row = DB.query(
            "SELECT * FROM users WHERE email=%s AND password_hash=%s AND status='ACT'",
            (email, DB.hash(pwd)), fetch='one'
        )
        
        if row:
            self.master.current_user = row
            self.master.show_dashboard()
        else:
            self.err_lbl.configure(text="Неверный email или пароль")


# Регистрация
class RegisterPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_CARD, corner_radius=12)
        self.master = master
        
        form = ctk.CTkFrame(self, fg_color="white", width=420, height=580, corner_radius=12)
        form.place(relx=0.5, rely=0.5, anchor="center")
        form.pack_propagate(False)
        
        ctk.CTkLabel(form, text="Регистрация", font=("Arial", 24, "bold")).pack(pady=(25, 15))
        
        self.email = ctk.CTkEntry(form, placeholder_text="Email", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.email.pack(pady=6)
        
        self.uname = ctk.CTkEntry(form, placeholder_text="Никнейм", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.uname.pack(pady=6)
        
        self.fname = ctk.CTkEntry(form, placeholder_text="Имя", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.fname.pack(pady=6)
        
        self.lname = ctk.CTkEntry(form, placeholder_text="Фамилия", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.lname.pack(pady=6)
        
        self.pwd = ctk.CTkEntry(form, placeholder_text="Пароль (мин. 6 символов)", show="*", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.pwd.pack(pady=6)
        
        self.pwd2 = ctk.CTkEntry(form, placeholder_text="Повторите пароль", show="*", width=340, height=40, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.pwd2.pack(pady=6)
        
        self.err_lbl = ctk.CTkLabel(form, text="", text_color="red")
        self.err_lbl.pack()
        
        ctk.CTkButton(form, text="Зарегистрироваться", fg_color=BG_CARD, text_color="black", 
                      width=340, height=45, corner_radius=8, command=self.handle_reg).pack(pady=15)
        ctk.CTkButton(form, text="← Назад к авторизации", fg_color="transparent", text_color="black", 
                      width=200, command=self.master.show_login).pack(pady=5)

    def handle_reg(self):
        email = self.email.get()
        uname = self.uname.get()
        fname = self.fname.get()
        lname = self.lname.get()
        pwd = self.pwd.get()
        pwd2 = self.pwd2.get()
        
        if not all([email, uname, fname, lname, pwd, pwd2]):
            self.err_lbl.configure(text="Заполните все поля")
            return
        if pwd != pwd2:
            self.err_lbl.configure(text="Пароли не совпадают")
            return
        if len(pwd) < 6:
            self.err_lbl.configure(text="Пароль минимум 6 символов")
            return
        
        try:
            DB.query(
                "INSERT INTO users (email, password_hash, username, first_name, last_name, role, status)"
                " VALUES (%s,%s,%s,%s,%s,'TUR','ACT')",
                (email, DB.hash(pwd), uname, fname, lname), fetch=None
            )
            messagebox.showinfo("Успех", "Регистрация успешна! Войдите в систему.")
            self.master.show_login()
        except:
            self.err_lbl.configure(text="Email или никнейм уже используется")


# Восстановление пароля
class ForgotPasswordPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_CARD, corner_radius=12)
        self.master = master
        
        form = ctk.CTkFrame(self, fg_color="white", width=400, height=320, corner_radius=12)
        form.place(relx=0.5, rely=0.5, anchor="center")
        form.pack_propagate(False)
        
        ctk.CTkLabel(form, text="Восстановление пароля", font=("Arial", 20, "bold")).pack(pady=(30, 20))
        
        self.email = ctk.CTkEntry(form, placeholder_text="Укажите ваш Email", width=320, height=45, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
        self.email.pack(pady=15)
        self.msg = ctk.CTkLabel(form, text="", text_color="gray")
        self.msg.pack()
        
        ctk.CTkButton(form, text="Сбросить пароль", fg_color=BG_CARD, text_color="black", 
                      width=320, height=45, corner_radius=8, command=self.show_popup).pack(pady=15)
        ctk.CTkButton(form, text="← Войти", fg_color="transparent", text_color="black", 
                      width=120, command=self.master.show_login).pack()

    def show_popup(self):
        email = self.email.get()
        if not email:
            self.msg.configure(text="Введите email")
            return
        
        row = DB.query("SELECT id FROM users WHERE email=%s", (email,), fetch='one')
        if not row:
            self.msg.configure(text="Email не найден в системе")
            return
        
        popup = ctk.CTkToplevel(self)
        popup.title("Статус")
        popup.geometry("320x150")
        popup.configure(fg_color=BG_WHITE)
        popup.transient(self)
        popup.grab_set()
        
        ctk.CTkLabel(popup, text="✅ Ссылка успешно отправлена!", font=("Arial", 14, "bold"), text_color="green").pack(pady=30)
        ctk.CTkButton(popup, text="Отлично", fg_color=BG_CARD, text_color="black", width=150, height=35,
                      command=lambda: [popup.destroy(), self.master.show_login()]).pack()


# Дашборд админа
class AdminDashboardPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=4)
        self.grid_rowconfigure(0, weight=1)
        
        # Левая панель меню
        sidebar = ctk.CTkFrame(self, fg_color=BG_CARD, corner_radius=12)
        sidebar.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        
        ctk.CTkLabel(sidebar, text="Guide-Advisor\nПанель админа", font=("Arial", 18, "bold")).pack(pady=25)
        
        menus = [
            ("🏠 Дашборд", lambda: master.show_dashboard()),
            ("👤 Каталог гидов", master.show_guides_catalog),
            ("🗺️ Каталог экскурсий", master.show_excursions_catalog),
            ("⚙️ Управление данными", master.show_manage_data),
            ("➕ Добавить гида", master.show_add_guide),
            ("➕ Добавить экскурсию", master.show_add_excursion),
            ("🔎 Глобальный поиск", master.show_search_catalog),
            ("📊 Отчеты", master.show_reports),
            ("🚪 Выход", self.logout),
        ]
        
        for name, cmd in menus:
            ctk.CTkButton(sidebar, text=name, fg_color="white", text_color="black", 
                          hover_color=BG_MAIN, anchor="center", height=40, width=180, corner_radius=8,
                          command=cmd).pack(padx=12, pady=5)
        
        self.content = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        self.content.grid(row=0, column=1, sticky="nsew")
        self._load_dashboard()

    def _load_dashboard(self):
        for w in self.content.winfo_children():
            w.destroy()
        
        total_guides = (DB.query("SELECT COUNT(*) as c FROM guides", fetch='one') or {}).get('c', 0)
        total_exc = (DB.query("SELECT COUNT(*) as c FROM excursions WHERE status='PUB'", fetch='one') or {}).get('c', 0)
        total_users = (DB.query("SELECT COUNT(*) as c FROM users", fetch='one') or {}).get('c', 0)
        total_bookings = (DB.query("SELECT COUNT(*) as c FROM bookings", fetch='one') or {}).get('c', 0)
        
        # Увеличенная панель статистики
        stats_layout = ctk.CTkFrame(self.content, fg_color="transparent")
        stats_layout.pack(fill="x", padx=20, pady=20)
        
        stats_data = [
            ("👤 Гидов", total_guides, "#E8E0D5"),
            ("🗺️ Экскурсий", total_exc, "#DDD5C8"),
            ("👥 Пользователей", total_users, "#D4CBB8"),
            ("📅 Бронирований", total_bookings, "#CEC4B0"),
        ]
        
        for label, value, color in stats_data:
            box = ctk.CTkFrame(stats_layout, fg_color=color, height=160, corner_radius=12)
            box.pack(side="left", expand=True, fill="x", padx=8)
            
            ctk.CTkLabel(box, text=label, font=("Arial", 16, "bold"), text_color="#4A4A4A").pack(pady=(20, 10))
            ctk.CTkLabel(box, text=str(value), font=("Arial", 56, "bold"), text_color="#2C2C2C").pack(pady=10)
        
        # Гистограмма
        ctk.CTkLabel(self.content, text="📈 Активность бронирований по дням", font=("Arial", 16, "bold")).pack(anchor="w", padx=20, pady=(20, 10))
        
        graph_frame = ctk.CTkFrame(self.content, fg_color=BG_MAIN, height=200, corner_radius=12)
        graph_frame.pack(fill="x", padx=20, pady=10)
        
        booking_stats = DB.query(
            """SELECT DATE(booked_at) as date, COUNT(*) as count 
               FROM bookings 
               WHERE booked_at > DATE_SUB(NOW(), INTERVAL 7 DAY)
               GROUP BY DATE(booked_at)
               ORDER BY date ASC""", fetch='all'
        )
        
        days_data = {}
        for i in range(6, -1, -1):
            date = datetime.now() - timedelta(days=i)
            day_name = date.strftime('%a')
            days_data[day_name] = 0
        
        for stat in booking_stats:
            if stat['date']:
                day_name = stat['date'].strftime('%a')
                if day_name in days_data:
                    days_data[day_name] = stat['count']
        
        max_val = max(days_data.values()) if days_data.values() else 1
        if max_val == 0:
            max_val = 1
        
        bars_frame = ctk.CTkFrame(graph_frame, fg_color="transparent")
        bars_frame.pack(fill="both", expand=True, padx=20, pady=15)
        
        for day, val in days_data.items():
            bar_container = ctk.CTkFrame(bars_frame, fg_color="transparent")
            bar_container.pack(side="left", expand=True, fill="y", padx=5)
            
            height_percent = (val / max_val) * 120
            bar_height = max(20, height_percent)
            
            bar = ctk.CTkFrame(bar_container, fg_color=BG_CARD, width=50, height=bar_height, corner_radius=6)
            bar.pack(side="bottom", pady=(5, 0))
            
            ctk.CTkLabel(bar_container, text=str(val), font=("Arial", 11, "bold")).pack(side="bottom")
            ctk.CTkLabel(bar_container, text=day, font=("Arial", 11)).pack(side="bottom", pady=(5, 0))
        
        # Список экскурсий
        ctk.CTkLabel(self.content, text="📋 Последние добавленные экскурсии", font=("Arial", 16, "bold")).pack(anchor="w", padx=20, pady=(20, 10))
        
        rows = DB.query(
            """SELECT e.id, e.title, e.category, e.status, u.first_name, u.last_name
               FROM excursions e LEFT JOIN users u ON e.guide_id = u.id
               ORDER BY e.id DESC"""
        )
        
        exc_scroll = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        exc_scroll.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        
        for r in rows:
            card = ctk.CTkFrame(exc_scroll, fg_color=BG_MAIN, corner_radius=10)
            card.pack(fill="x", pady=6)
            
            status_text = "🟢 Активна" if r['status'] == 'PUB' else "🟡 На модерации"
            
            ctk.CTkLabel(card, text=f"{r['title']}", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=(10, 3))
            ctk.CTkLabel(card, text=f"Категория: {r['category']}  |  Гид: {r['first_name'] or ''} {r['last_name'] or ''}  |  {status_text}",
                        font=("Arial", 12)).pack(anchor="w", padx=15, pady=(0, 10))

    def _delete_excursion(self, exc_id, title):
        if not messagebox.askyesno('Удалить экскурсию', f'Удалить экскурсию "{title}"?\nВсе бронирования и отзывы также будут удалены.'):
            return
        DB.query("DELETE FROM reviews WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM bookings WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM schedules WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM excursions WHERE id=%s", (exc_id,), fetch=None)
        messagebox.showinfo("Готово", "Экскурсия удалена.")
        self._load_dashboard()

    def logout(self):
        self.master.current_user = None
        self.master.show_login()


# Дашборд гида
class GuideDashboardPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=4)
        self.grid_rowconfigure(0, weight=1)
        
        sidebar = ctk.CTkFrame(self, fg_color=BG_CARD, corner_radius=12)
        sidebar.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        
        ctk.CTkLabel(sidebar, text="Guide-Advisor\nПанель гида", font=("Arial", 18, "bold")).pack(pady=25)
        
        menus = [
            ("🏠 Дашборд", lambda: master.show_dashboard()),
            ("👤 Каталог гидов", master.show_guides_catalog),
            ("🗺️ Каталог экскурсий", master.show_excursions_catalog),
            ("⭐ Мои отзывы", lambda: master.show_guide_reviews(all_reviews=False)),
            ("📝 Все отзывы", lambda: master.show_guide_reviews(all_reviews=True)),
            ("📅 Мое расписание", master.show_schedule),
            ("➕ Добавить экскурсию", master.show_add_excursion),
            ("🔔 Уведомления", master.show_notifications),
            ("🔎 Глобальный поиск", master.show_search_catalog),
            ("🚪 Выход", self.logout),
        ]
        
        for name, cmd in menus:
            ctk.CTkButton(sidebar, text=name, fg_color="white", text_color="black", 
                          hover_color=BG_MAIN, anchor="center", height=40, width=180, corner_radius=8,
                          command=cmd).pack(padx=12, pady=5)
        
        self.content = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        self.content.grid(row=0, column=1, sticky="nsew")
        self._load_dashboard()

    def _load_dashboard(self):
        for w in self.content.winfo_children():
            w.destroy()
        
        my_exc = (DB.query("SELECT COUNT(*) as c FROM excursions WHERE guide_id=%s", 
                          (self.master.current_user['id'],), fetch='one') or {}).get('c', 0)
        my_bookings = (DB.query("""
            SELECT COUNT(*) as c FROM bookings b
            JOIN schedules s ON b.schedule_id = s.id
            JOIN excursions e ON s.excursion_id = e.id
            WHERE e.guide_id=%s""", (self.master.current_user['id'],), fetch='one') or {}).get('c', 0)
        my_reviews = (DB.query("""
            SELECT COUNT(*) as c FROM reviews r
            JOIN excursions e ON r.excursion_id = e.id
            WHERE e.guide_id=%s""", (self.master.current_user['id'],), fetch='one') or {}).get('c', 0)
        unread_notif = (DB.query("SELECT COUNT(*) as c FROM notifications WHERE user_id=%s AND is_read=FALSE",
                                (self.master.current_user['id'],), fetch='one') or {}).get('c', 0)
        
        stats_layout = ctk.CTkFrame(self.content, fg_color="transparent")
        stats_layout.pack(fill="x", padx=20, pady=20)
        
        stats_data = [
            ("📋 Мои экскурсии", my_exc, "#E8E0D5"),
            ("📅 Бронирований", my_bookings, "#DDD5C8"),
            ("⭐ Отзывов", my_reviews, "#D4CBB8"),
            ("🔔 Уведомления", unread_notif, "#CEC4B0"),
        ]
        
        for label, value, color in stats_data:
            box = ctk.CTkFrame(stats_layout, fg_color=color, height=140, corner_radius=12)
            box.pack(side="left", expand=True, fill="x", padx=8)
            ctk.CTkLabel(box, text=label, font=("Arial", 14, "bold"), text_color="#4A4A4A").pack(pady=(20, 10))
            ctk.CTkLabel(box, text=str(value), font=("Arial", 48, "bold"), text_color="#2C2C2C").pack(pady=10)
        
        ctk.CTkLabel(self.content, text="Мои экскурсии", font=("Arial", 16, "bold")).pack(anchor="w", padx=20, pady=(20, 10))
        
        rows = DB.query(
            "SELECT id, title, category, status FROM excursions WHERE guide_id=%s ORDER BY id DESC",
            (self.master.current_user['id'],)
        )
        
        exc_scroll = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        exc_scroll.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        
        for r in rows:
            card = ctk.CTkFrame(exc_scroll, fg_color=BG_MAIN, corner_radius=10)
            card.pack(fill="x", pady=6)
            
            btn_frame = ctk.CTkFrame(card, fg_color="transparent")
            btn_frame.pack(side="right", padx=15, pady=8)
            ctk.CTkButton(btn_frame, text="✏️ Редактировать", width=130, height=35, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda eid=r['id']: self.master.show_edit_excursion(eid)).pack()
            ctk.CTkButton(btn_frame, text="⭐ Отзывы", width=130, height=35, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda eid=r['id']: self.master.show_reviews(eid)).pack(pady=(5, 0))
            ctk.CTkButton(btn_frame, text="🗑️ Удалить", width=130, height=35, fg_color="#E57373", text_color="white", corner_radius=6,
                         hover_color="#C62828",
                         command=lambda eid=r['id'], t=r['title']: self._delete_excursion(eid, t)).pack(pady=(5, 0))
            
            status_text = "🟢 Активна" if r['status'] == 'PUB' else "🔴 Скрыта"
            ctk.CTkLabel(card, text=f"{r['title']} | {r['category']} | {status_text}",
                        font=("Arial", 13)).pack(anchor="w", padx=15, pady=12)

    def _delete_excursion(self, exc_id, title):
        if not messagebox.askyesno('Удалить', f'Удалить экскурсию "{title}"?\nВсе отзывы и бронирования также будут удалены.'):
            return
        DB.query("DELETE FROM reviews WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM bookings WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM schedules WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM excursions WHERE id=%s", (exc_id,), fetch=None)
        messagebox.showinfo("Готово", "Экскурсия удалена.")
        self.master.show_dashboard()

    def logout(self):
        self.master.current_user = None
        self.master.show_login()


# Страница отзывов для гида
class GuideReviewsPage(ctk.CTkFrame):
    def __init__(self, master, all_reviews=False):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        self.all_reviews = all_reviews
        
        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black", 
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        title = "Все отзывы на экскурсии" if all_reviews else "Отзывы на мои экскурсии"
        ctk.CTkLabel(self, text=title, font=("Arial", 22, "bold")).pack(pady=15)
        
        scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=20, pady=10)
        
        if all_reviews:
            reviews = DB.query(
                """SELECT r.*, u.first_name as user_first_name, u.last_name as user_last_name, 
                          e.title as excursion_title, e.guide_id, gu.first_name as guide_first_name, gu.last_name as guide_last_name
                   FROM reviews r
                   JOIN excursions e ON r.excursion_id = e.id
                   JOIN users u ON r.user_id = u.id
                   JOIN users gu ON e.guide_id = gu.id
                   ORDER BY r.created_at DESC""",
                fetch='all'
            )
        else:
            reviews = DB.query(
                """SELECT r.*, u.first_name, u.last_name, e.title as excursion_title, e.guide_id
                   FROM reviews r
                   JOIN excursions e ON r.excursion_id = e.id
                   JOIN users u ON r.user_id = u.id
                   WHERE e.guide_id = %s
                   ORDER BY r.created_at DESC""",
                (self.master.current_user['id'],), fetch='all'
            )
        
        if not reviews:
            ctk.CTkLabel(scroll, text="Пока нет отзывов", font=("Arial", 14), text_color="gray").pack(pady=30)
            return
        
        total_reviews = len(reviews)
        avg_rating = sum(r['rating'] for r in reviews) / total_reviews if total_reviews > 0 else 0
        
        stats_frame = ctk.CTkFrame(scroll, fg_color="transparent", corner_radius=12)
        stats_frame.pack(fill="x", pady=10)
        
        ctk.CTkLabel(stats_frame, text="📈 Статистика отзывов", font=("Arial", 14, "bold")).pack(anchor="w", padx=0, pady=10)
        
        stats_grid = ctk.CTkFrame(stats_frame, fg_color="transparent")
        stats_grid.pack(fill="x", pady=5)
        stats_grid.grid_columnconfigure(0, weight=1)
        stats_grid.grid_columnconfigure(1, weight=1)
        
        box1 = ctk.CTkFrame(stats_grid, fg_color="white", corner_radius=12, height=120)
        box1.grid(row=0, column=0, padx=(0, 8), pady=5, sticky="nsew")
        box1.pack_propagate(False)
        ctk.CTkLabel(box1, text="Всего отзывов", font=("Arial", 14, "bold"), text_color="#4A4A4A").pack(pady=(20, 5))
        ctk.CTkLabel(box1, text=str(total_reviews), font=("Arial", 40, "bold"), text_color="#2C2C2C").pack()
        
        box2 = ctk.CTkFrame(stats_grid, fg_color="white", corner_radius=12, height=120)
        box2.grid(row=0, column=1, padx=(8, 0), pady=5, sticky="nsew")
        box2.pack_propagate(False)
        ctk.CTkLabel(box2, text="Средний рейтинг", font=("Arial", 14, "bold"), text_color="#4A4A4A").pack(pady=(20, 5))
        ctk.CTkLabel(box2, text=f"{avg_rating:.1f} ★", font=("Arial", 40, "bold"), text_color="#2C2C2C").pack()
        
        for r in reviews:
            card = ctk.CTkFrame(scroll, fg_color="white", corner_radius=12)
            card.pack(fill="x", pady=8)
            
            stars = "★" * r['rating'] + "☆" * (5 - r['rating'])
            
            if all_reviews:
                guide_name = f"{r['guide_first_name']} {r['guide_last_name']}"
                ctk.CTkLabel(card, text=f"📌 {r['excursion_title']}", font=("Arial", 15, "bold")).pack(anchor="w", padx=15, pady=(10, 5))
                ctk.CTkLabel(card, text=f"Гид: {guide_name} | От: {r['user_first_name']} {r['user_last_name']} • {r['created_at'].strftime('%d.%m.%Y')}",
                            font=("Arial", 11), text_color="#666").pack(anchor="w", padx=15)
            else:
                ctk.CTkLabel(card, text=f"📌 {r['excursion_title']}", font=("Arial", 15, "bold")).pack(anchor="w", padx=15, pady=(10, 5))
                ctk.CTkLabel(card, text=f"От: {r['first_name']} {r['last_name']} • {r['created_at'].strftime('%d.%m.%Y')}",
                            font=("Arial", 11), text_color="#666").pack(anchor="w", padx=15)
            
            ctk.CTkLabel(card, text=stars, font=("Arial", 12)).pack(anchor="w", padx=15, pady=5)
            ctk.CTkLabel(card, text=f"\"{r['comment']}\"", wraplength=700, justify="left", 
                        font=("Arial", 12), text_color="#333").pack(anchor="w", padx=15, pady=(0, 10))


# Дашборд туриста
class TouristDashboardPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_columnconfigure(1, weight=4)
        self.grid_rowconfigure(0, weight=1)
        
        sidebar = ctk.CTkFrame(self, fg_color=BG_CARD, corner_radius=12)
        sidebar.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        
        ctk.CTkLabel(sidebar, text="Guide-Advisor\nПанель туриста", font=("Arial", 18, "bold")).pack(pady=25)
        
        menus = [
            ("🏠 Дашборд", lambda: master.show_dashboard()),
            ("👤 Каталог гидов", master.show_guides_catalog),
            ("🗺️ Каталог экскурсий", master.show_excursions_catalog),
            ("⭐ Мои отзывы", master.show_reviews),
            ("📅 Мое расписание", master.show_schedule),
            ("🎯 Рекомендации", master.show_personal_recommendations),
            ("🔔 Уведомления", master.show_notifications),
            ("👤 Мой профиль", master.show_tourist_profile),
            ("🔎 Глобальный поиск", master.show_search_catalog),
            ("🚪 Выход", self.logout),
        ]
        
        for name, cmd in menus:
            ctk.CTkButton(sidebar, text=name, fg_color="white", text_color="black", 
                          hover_color=BG_MAIN, anchor="center", height=40, width=180, corner_radius=8,
                          command=cmd).pack(padx=12, pady=5)
        
        self.content = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        self.content.grid(row=0, column=1, sticky="nsew")
        self._load_dashboard()

    def _load_dashboard(self):
        for w in self.content.winfo_children():
            w.destroy()
        
        total_guides = (DB.query("SELECT COUNT(*) as c FROM guides", fetch='one') or {}).get('c', 0)
        total_exc = (DB.query("SELECT COUNT(*) as c FROM excursions WHERE status='PUB'", fetch='one') or {}).get('c', 0)
        my_bookings = (DB.query("SELECT COUNT(*) as c FROM bookings WHERE tourist_id=%s", 
                               (self.master.current_user['id'],), fetch='one') or {}).get('c', 0)
        
        unread_notif = DB.query("SELECT COUNT(*) as c FROM notifications WHERE user_id=%s AND is_read=FALSE",
                                (self.master.current_user['id'],), fetch='one') or {}
        unread_count = unread_notif.get('c', 0)
        
        stats_layout = ctk.CTkFrame(self.content, fg_color="transparent")
        stats_layout.pack(fill="x", padx=20, pady=20)
        
        stats_data = [
            ("👤 Гидов", total_guides, "#E8E0D5"),
            ("🗺️ Экскурсий", total_exc, "#DDD5C8"),
            ("📅 Мои брони", my_bookings, "#D4CBB8"),
            ("🔔 Уведомления", unread_count, "#CEC4B0"),
        ]
        
        for label, value, color in stats_data:
            box = ctk.CTkFrame(stats_layout, fg_color=color, height=140, corner_radius=12)
            box.pack(side="left", expand=True, fill="x", padx=8)
            ctk.CTkLabel(box, text=label, font=("Arial", 14, "bold"), text_color="#4A4A4A").pack(pady=(20, 10))
            ctk.CTkLabel(box, text=str(value), font=("Arial", 48, "bold"), text_color="#2C2C2C").pack(pady=10)
        
        ctk.CTkLabel(self.content, text="Доступные экскурсии", font=("Arial", 16, "bold")).pack(anchor="w", padx=20, pady=(20, 10))
        
        rows = DB.query(
            """SELECT e.id, e.title, e.category, e.duration_hours, e.max_participants, e.description,
                      u.first_name, u.last_name
               FROM excursions e LEFT JOIN users u ON e.guide_id = u.id
               WHERE e.status = 'PUB'
               ORDER BY e.id DESC"""
        )
        
        exc_scroll = ctk.CTkScrollableFrame(self.content, fg_color="transparent")
        exc_scroll.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        
        for r in rows:
            card = ctk.CTkFrame(exc_scroll, fg_color=BG_MAIN, corner_radius=10)
            card.pack(fill="x", pady=6)
            
            info = ctk.CTkFrame(card, fg_color="transparent")
            info.pack(side="left", fill="both", expand=True, padx=15, pady=10)
            
            ctk.CTkLabel(info, text=r['title'], font=("Arial", 15, "bold")).pack(anchor="w")
            ctk.CTkLabel(info, text=f"Гид: {r['first_name']} {r['last_name']}  •  {r['category']}", font=("Arial", 12)).pack(anchor="w")
            ctk.CTkLabel(info, text=f"Длительность: {r['duration_hours']} ч  •  Макс: {r['max_participants']} чел.", font=("Arial", 12)).pack(anchor="w")
            
            btn_frame = ctk.CTkFrame(card, fg_color="transparent")
            btn_frame.pack(side="right", padx=15, pady=10)
            ctk.CTkButton(btn_frame, text="Забронировать", width=140, height=35, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda eid=r['id']: self._book_excursion(eid)).pack()
            ctk.CTkButton(btn_frame, text="⭐ Отзывы", width=140, height=35, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda eid=r['id']: self.master.show_reviews(eid)).pack(pady=(5, 0))

    def _book_excursion(self, exc_id):
        e = DB.query("SELECT * FROM excursions WHERE id=%s", (exc_id,), fetch='one')
        if not e:
            messagebox.showerror("Ошибка", "Экскурсия не найдена")
            return
        
        self.master.log_activity(self.master.current_user['id'], exc_id, 'VIEW')
        
        schedules = DB.query(
            "SELECT id, start_datetime FROM schedules WHERE excursion_id=%s AND start_datetime > NOW()",
            (exc_id,)
        )
        
        if not schedules:
            messagebox.showerror("Ошибка", "Нет доступных дат")
            return
        
        win = ctk.CTkToplevel(self.master)
        win.title("Бронирование")
        win.geometry("420x420")
        win.configure(fg_color=BG_WHITE)
        win.transient(self.master)
        win.grab_set()
        
        ctk.CTkLabel(win, text=e['title'], font=("Arial", 18, "bold")).pack(pady=20)
        ctk.CTkLabel(win, text="Выберите дату:", font=("Arial", 13)).pack()
        
        schedule_var = ctk.StringVar(value=schedules[0]['start_datetime'].strftime('%d.%m.%Y %H:%M'))
        schedule_combo = ctk.CTkComboBox(win, variable=schedule_var,
                                        values=[s['start_datetime'].strftime('%d.%m.%Y %H:%M') for s in schedules],
                                        fg_color=BG_MAIN, width=280, height=35, corner_radius=8)
        schedule_combo.pack(pady=10)
        
        ctk.CTkLabel(win, text="Количество участников:", font=("Arial", 13)).pack()
        parts_var = ctk.IntVar(value=1)
        parts_spin = ctk.CTkEntry(win, textvariable=parts_var, width=100, height=35, fg_color=BG_MAIN, corner_radius=8)
        parts_spin.pack(pady=5)
        
        def confirm_booking():
            schedule_id = None
            for s in schedules:
                if s['start_datetime'].strftime('%d.%m.%Y %H:%M') == schedule_var.get():
                    schedule_id = s['id']
                    break
            if schedule_id:
                DB.query(
                    "INSERT INTO bookings (tourist_id, schedule_id, participants_count, status)"
                    " VALUES (%s,%s,%s,'PEND')",
                    (self.master.current_user['id'], schedule_id, parts_var.get()), fetch=None
                )
                self.master.log_activity(self.master.current_user['id'], exc_id, 'BOOK')
                self.master.add_notification(self.master.current_user['id'], 'Бронирование создано',
                                            f'Вы успешно забронировали экскурсию "{e["title"]}"', 'BOOKING')
                # Уведомление гиду о новом бронировании
                user = self.master.current_user
                self.master.add_notification(
                    e['guide_id'],
                    'Новое бронирование',
                    f'{user["first_name"]} {user["last_name"]} забронировал вашу экскурсию "{e["title"]}"',
                    'BOOKING'
                )
                messagebox.showinfo("Успех", "Бронирование создано!")
                win.destroy()
                self._load_dashboard()
            else:
                messagebox.showerror("Ошибка", "Выберите дату")
        
        ctk.CTkButton(win, text="Подтвердить", width=200, height=40, fg_color=BG_CARD, text_color="black", corner_radius=8,
                     command=confirm_booking).pack(pady=25)

    def logout(self):
        self.master.current_user = None
        self.master.show_login()


# Страница отзывов
class ReviewsPage(ctk.CTkFrame):
    def __init__(self, master, excursion_id=None):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        self.excursion_id = excursion_id

        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black",
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)

        ctk.CTkLabel(self, text="Отзывы об экскурсиях", font=("Arial", 22, "bold")).pack(pady=15)

        if excursion_id and self.master.current_user['role'] == 'TUR':
            self._add_review_form()

        self.scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.scroll.pack(fill="both", expand=True, padx=20, pady=10)

        self._load_reviews()
    
    def _add_review_form(self):
        form_card = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        form_card.pack(fill="x", padx=20, pady=10)
        
        ctk.CTkLabel(form_card, text="Оставить отзыв", font=("Arial", 16, "bold")).pack(pady=10)
        
        exc = DB.query("SELECT title FROM excursions WHERE id=%s", (self.excursion_id,), fetch='one')
        if exc:
            ctk.CTkLabel(form_card, text=f"Экскурсия: {exc['title']}", font=("Arial", 12)).pack()
        
        rating_frame = ctk.CTkFrame(form_card, fg_color="transparent")
        rating_frame.pack(pady=5)
        ctk.CTkLabel(rating_frame, text="Оценка:").pack(side="left")
        self.rating_var = ctk.IntVar(value=5)
        for i in range(1, 6):
            ctk.CTkRadioButton(rating_frame, text=str(i), variable=self.rating_var, value=i).pack(side="left", padx=5)
        
        self.comment_entry = ctk.CTkTextbox(form_card, width=500, height=80, fg_color=BG_MAIN, corner_radius=8)
        self.comment_entry.pack(pady=10)
        
        ctk.CTkButton(form_card, text="Отправить отзыв", fg_color=BG_CARD, text_color="black", 
                      width=200, height=35, corner_radius=8, command=self._submit_review).pack(pady=10)
    
    def _submit_review(self):
        if not self.excursion_id:
            return
        
        rating = self.rating_var.get()
        comment = self.comment_entry.get("1.0", "end-1c")
        
        if not comment.strip():
            messagebox.showerror("Ошибка", "Введите текст отзыва")
            return
        
        DB.query(
            "INSERT INTO reviews (excursion_id, user_id, rating, comment) VALUES (%s, %s, %s, %s)",
            (self.excursion_id, self.master.current_user['id'], rating, comment), fetch=None
        )
        
        self.master.log_activity(self.master.current_user['id'], self.excursion_id, 'REVIEW')
        
        # Уведомление туристу об успешном отзыве
        exc = DB.query("SELECT title, guide_id FROM excursions WHERE id=%s", (self.excursion_id,), fetch='one')
        if exc:
            self.master.add_notification(
                self.master.current_user['id'],
                'Отзыв опубликован',
                f'Ваш отзыв на экскурсию "{exc["title"]}" успешно добавлен',
                'INFO'
            )
            # Уведомление гиду о новом отзыве
            user = self.master.current_user
            self.master.add_notification(
                exc['guide_id'],
                'Новый отзыв',
                f'{user["first_name"]} {user["last_name"]} оставил отзыв на вашу экскурсию "{exc["title"]}" — {rating}★',
                'INFO'
            )
        
        messagebox.showinfo("Успех", "Спасибо за отзыв!")
        self._load_reviews()
        self.comment_entry.delete("1.0", "end")
    
    def _load_reviews(self):
        for w in self.scroll.winfo_children():
            w.destroy()

        if self.excursion_id:
            reviews = DB.query(
                """SELECT r.*, u.first_name, u.last_name 
                   FROM reviews r 
                   JOIN users u ON r.user_id = u.id 
                   WHERE r.excursion_id = %s 
                   ORDER BY r.created_at DESC""",
                (self.excursion_id,), fetch='all'
            )
            exc_info = DB.query("SELECT title FROM excursions WHERE id=%s", (self.excursion_id,), fetch='one')
            if exc_info:
                ctk.CTkLabel(self.scroll, text=f"Экскурсия: {exc_info['title']}", font=("Arial", 14, "bold")).pack(anchor="w", pady=10)
        else:
            reviews = DB.query(
                """SELECT r.*, u.first_name, u.last_name, e.title 
                   FROM reviews r 
                   JOIN users u ON r.user_id = u.id 
                   JOIN excursions e ON r.excursion_id = e.id 
                   ORDER BY r.created_at DESC""",
                fetch='all'
            )
        
        if not reviews:
            ctk.CTkLabel(self.scroll, text="Пока нет отзывов. Будьте первым!", font=("Arial", 12), text_color="gray").pack(pady=20)
            return
        
        current_user = self.master.current_user
        for r in reviews:
            card = ctk.CTkFrame(self.scroll, fg_color="white", corner_radius=10)
            card.pack(fill="x", pady=8)

            stars = "★" * r['rating'] + "☆" * (5 - r['rating'])
            title_text = f"{r['title']} - " if 'title' in r else ""

            inner = ctk.CTkFrame(card, fg_color="transparent")
            inner.pack(fill="x", padx=15, pady=10)

            top_row = ctk.CTkFrame(inner, fg_color="transparent")
            top_row.pack(fill="x")
            ctk.CTkLabel(top_row, text=f"{title_text}{r['first_name']} {r['last_name']} • {r['created_at'].strftime('%d.%m.%Y')}",
                        font=("Arial", 11, "bold"), text_color="#666").pack(side="left")
            if current_user and r['user_id'] == current_user['id']:
                ctk.CTkButton(top_row, text="🗑️ Удалить", width=110, height=30, fg_color="#E57373", text_color="white",
                             corner_radius=6, hover_color="#C62828",
                             command=lambda rid=r['id']: self._delete_review(rid)).pack(side="right")
            ctk.CTkLabel(inner, text=stars, font=("Arial", 12)).pack(anchor="w", pady=(4, 0))
            ctk.CTkLabel(inner, text=r['comment'], wraplength=700, justify="left", font=("Arial", 12)).pack(anchor="w", pady=(4, 0))


    def _delete_review(self, review_id):
        if not messagebox.askyesno("Удалить отзыв", "Удалить ваш отзыв?"):
            return
        DB.query("DELETE FROM reviews WHERE id=%s", (review_id,), fetch=None)
        self._load_reviews()


# Страница расписания
class SchedulePage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black", 
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        ctk.CTkLabel(self, text="Мое расписание", font=("Arial", 22, "bold")).pack(pady=15)
        
        filter_frame = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        filter_frame.pack(fill="x", padx=20, pady=10)

        ctk.CTkLabel(filter_frame, text="Фильтр:").pack(side="left", padx=15)

        self.filter_var = ctk.StringVar(value="all")
        ctk.CTkRadioButton(filter_frame, text="Все", variable=self.filter_var, value="all", command=self._load_schedule).pack(side="left", padx=10)
        ctk.CTkRadioButton(filter_frame, text="Предстоящие", variable=self.filter_var, value="upcoming", command=self._load_schedule).pack(side="left", padx=10)
        ctk.CTkRadioButton(filter_frame, text="Прошедшие", variable=self.filter_var, value="past", command=self._load_schedule).pack(side="left", padx=10)

        if master.current_user['role'] == 'GUI':
            ctk.CTkButton(filter_frame, text="➕ Добавить дату", fg_color=BG_CARD, text_color="black",
                         width=120, height=18, corner_radius=8, command=self._add_schedule_dialog).pack(side="right", padx=15)

        self.scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.scroll.pack(fill="both", expand=True, padx=20, pady=10)
        
        self._load_schedule()
    
    def _load_schedule(self):
        for w in self.scroll.winfo_children():
            w.destroy()
        
        user = self.master.current_user
        
        if user['role'] == 'TUR':
            bookings = DB.query(
                """SELECT b.*, e.title, e.id as excursion_id, s.start_datetime, u.first_name, u.last_name
                   FROM bookings b
                   JOIN schedules s ON b.schedule_id = s.id
                   JOIN excursions e ON s.excursion_id = e.id
                   JOIN users u ON e.guide_id = u.id
                   WHERE b.tourist_id = %s
                   ORDER BY s.start_datetime ASC""",
                (user['id'],), fetch='all'
            )
            
            for b in bookings:
                self._create_tourist_schedule_card(b)
        else:
            excursions = DB.query(
                """SELECT e.*, s.id as schedule_id, s.start_datetime,
                          (SELECT COUNT(*) FROM bookings b WHERE b.schedule_id = s.id) as bookings_count
                   FROM excursions e
                   JOIN schedules s ON e.id = s.excursion_id
                   WHERE e.guide_id = %s
                   ORDER BY s.start_datetime ASC""",
                (user['id'],), fetch='all'
            )
            
            for e in excursions:
                self._create_guide_schedule_card(e)
    
    def _create_tourist_schedule_card(self, booking):
        now = datetime.now()
        is_upcoming = booking['start_datetime'] > now
        
        if self.filter_var.get() == "upcoming" and not is_upcoming:
            return
        if self.filter_var.get() == "past" and is_upcoming:
            return
        
        status_colors = {
            'PEND': "#FF9800",
            'CONF': "#4CAF50",
            'CANC': "#F44336",
            'COM': "#9E9E9E"
        }
        status_texts = {
            'PEND': "⏳ Ожидает",
            'CONF': "✅ Подтверждено",
            'CANC': "❌ Отменено",
            'COM': "📌 Завершено"
        }
        
        card = ctk.CTkFrame(self.scroll, fg_color="white", corner_radius=12)
        card.pack(fill="x", pady=8)
        
        color = status_colors.get(booking['status'], "#666")
        date_str = booking['start_datetime'].strftime('%d.%m.%Y %H:%M')
        
        ctk.CTkLabel(card, text=booking['title'], font=("Arial", 16, "bold")).pack(anchor="w", padx=15, pady=(10, 5))
        ctk.CTkLabel(card, text=f"Гид: {booking['first_name']} {booking['last_name']}", font=("Arial", 12)).pack(anchor="w", padx=15)
        ctk.CTkLabel(card, text=f"Дата и время: {date_str}", font=("Arial", 12)).pack(anchor="w", padx=15)
        ctk.CTkLabel(card, text=f"Участников: {booking['participants_count']}", font=("Arial", 12)).pack(anchor="w", padx=15)
        ctk.CTkLabel(card, text=status_texts.get(booking['status'], booking['status']), 
                     font=("Arial", 12, "bold"), text_color=color).pack(anchor="w", padx=15, pady=(0, 10))
        
        if booking['status'] == 'CONF' and is_upcoming:
            btn_frame = ctk.CTkFrame(card, fg_color="transparent")
            btn_frame.pack(anchor="e", padx=15, pady=(0, 10))
            ctk.CTkButton(btn_frame, text="✏️ Оставить отзыв", width=140, height=32, fg_color=BG_CARD, 
                         text_color="black", corner_radius=6,
                         command=lambda eid=booking['excursion_id']: self.master.show_reviews(eid)).pack()
    
    def _create_guide_schedule_card(self, excursion):
        now = datetime.now()
        is_upcoming = excursion['start_datetime'] > now
        
        if self.filter_var.get() == "upcoming" and not is_upcoming:
            return
        if self.filter_var.get() == "past" and is_upcoming:
            return
        
        card = ctk.CTkFrame(self.scroll, fg_color="white", corner_radius=12)
        card.pack(fill="x", pady=8)
        
        date_str = excursion['start_datetime'].strftime('%d.%m.%Y %H:%M')
        
        ctk.CTkLabel(card, text=excursion['title'], font=("Arial", 16, "bold")).pack(anchor="w", padx=15, pady=(10, 5))
        ctk.CTkLabel(card, text=f"Дата и время: {date_str}", font=("Arial", 12)).pack(anchor="w", padx=15)
        ctk.CTkLabel(card, text=f"Записей: {excursion['bookings_count']} / {excursion['max_participants']}", 
                     font=("Arial", 12)).pack(anchor="w", padx=15, pady=(0, 10))
        
        if excursion['bookings_count'] > 0:
            ctk.CTkButton(card, text="📋 Список участников", width=160, height=32, fg_color=BG_CARD, 
                         text_color="black", corner_radius=6,
                         command=lambda eid=excursion['id'], sd=excursion['start_datetime']: self._show_participants(eid, sd)).pack(anchor="e", padx=15, pady=(0, 10))
    
    def _add_schedule_dialog(self):
        excursions = DB.query(
            "SELECT id, title FROM excursions WHERE guide_id=%s AND status='PUB' ORDER BY title",
            (self.master.current_user['id'],), fetch='all'
        )
        if not excursions:
            messagebox.showinfo("Нет экскурсий", "Сначала создайте экскурсию.")
            return

        win = ctk.CTkToplevel(self.master)
        win.title("Добавить дату")
        win.geometry("420x340")
        win.configure(fg_color="white")
        win.transient(self.master)
        win.grab_set()

        ctk.CTkLabel(win, text="Добавить дату проведения", font=("Arial", 16, "bold")).pack(pady=20)

        ctk.CTkLabel(win, text="Экскурсия:", font=("Arial", 13)).pack()
        exc_var = ctk.StringVar(value=excursions[0]['title'])
        ctk.CTkComboBox(win, variable=exc_var, values=[e['title'] for e in excursions],
                       width=300, height=35, fg_color=BG_MAIN, corner_radius=8).pack(pady=8)

        ctk.CTkLabel(win, text="Дата и время (ДД.ММ.ГГГГ ЧЧ:ММ):", font=("Arial", 13)).pack()
        date_entry = ctk.CTkEntry(win, width=300, height=35, fg_color=BG_MAIN, corner_radius=8,
                                  placeholder_text="например: 25.06.2025 10:00")
        date_entry.pack(pady=8)

        def save():
            exc_id = next((e['id'] for e in excursions if e['title'] == exc_var.get()), None)
            if not exc_id:
                messagebox.showerror("Ошибка", "Выберите экскурсию")
                return
            try:
                from datetime import datetime as dt
                date = dt.strptime(date_entry.get().strip(), '%d.%m.%Y %H:%M')
            except ValueError:
                messagebox.showerror("Ошибка", "Неверный формат даты. Используйте ДД.ММ.ГГГГ ЧЧ:ММ")
                return
            DB.query("INSERT INTO schedules (excursion_id, start_datetime) VALUES (%s, %s)",
                     (exc_id, date), fetch=None)
            messagebox.showinfo("Готово", "Дата добавлена!")
            win.destroy()
            self._load_schedule()

        ctk.CTkButton(win, text="Сохранить", fg_color=BG_CARD, text_color="black",
                     width=200, height=40, corner_radius=8, command=save).pack(pady=20)

    def _show_participants(self, excursion_id, start_datetime):
        bookings = DB.query(
            """SELECT b.*, u.first_name, u.last_name, u.email
               FROM bookings b
               JOIN users u ON b.tourist_id = u.id
               JOIN schedules s ON b.schedule_id = s.id
               WHERE s.excursion_id = %s AND s.start_datetime = %s""",
            (excursion_id, start_datetime), fetch='all'
        )
        
        win = ctk.CTkToplevel(self.master)
        win.title("Участники экскурсии")
        win.geometry("500x400")
        win.configure(fg_color=BG_WHITE)
        
        ctk.CTkLabel(win, text=f"Участники ({len(bookings)})", font=("Arial", 16, "bold")).pack(pady=15)
        
        scroll = ctk.CTkScrollableFrame(win, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=20, pady=10)
        
        for b in bookings:
            card = ctk.CTkFrame(scroll, fg_color=BG_MAIN, corner_radius=8)
            card.pack(fill="x", pady=5)
            ctk.CTkLabel(card, text=f"{b['first_name']} {b['last_name']}", font=("Arial", 13, "bold")).pack(anchor="w", padx=10, pady=5)
            ctk.CTkLabel(card, text=f"Email: {b['email']}", font=("Arial", 11)).pack(anchor="w", padx=10)
            ctk.CTkLabel(card, text=f"Участников: {b['participants_count']}", font=("Arial", 11)).pack(anchor="w", padx=10, pady=(0, 5))


# Страница персональных рекомендаций
class PersonalRecommendationsPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black", 
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        ctk.CTkLabel(self, text="Персональные рекомендации", font=("Arial", 22, "bold")).pack(pady=15)
        
        self._load_recommendations()
    
    def _load_recommendations(self):
        scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=20, pady=10)
        
        user_id = self.master.current_user['id']
        
        user_categories = DB.query(
            """SELECT DISTINCT e.category, COUNT(*) as count
               FROM bookings b
               JOIN schedules s ON b.schedule_id = s.id
               JOIN excursions e ON s.excursion_id = e.id
               WHERE b.tourist_id = %s AND b.status = 'CONF'
               GROUP BY e.category
               ORDER BY count DESC""",
            (user_id,), fetch='all'
        )
        
        viewed_categories = DB.query(
            """SELECT DISTINCT e.category, COUNT(*) as count
               FROM user_activity ua
               JOIN excursions e ON ua.excursion_id = e.id
               WHERE ua.user_id = %s AND ua.action_type = 'VIEW'
               GROUP BY e.category
               ORDER BY count DESC""",
            (user_id,), fetch='all'
        )
        
        categories_interest = {}
        for cat in user_categories:
            categories_interest[cat['category']] = categories_interest.get(cat['category'], 0) + cat['count']
        for cat in viewed_categories:
            categories_interest[cat['category']] = categories_interest.get(cat['category'], 0) + cat['count']
        
        if categories_interest:
            top_category = max(categories_interest, key=categories_interest.get)
            ctk.CTkLabel(scroll, text=f"🎯 На основе ваших предпочтений", font=("Arial", 16, "bold")).pack(anchor="w", pady=10)
            ctk.CTkLabel(scroll, text=f"Ваш любимый тип экскурсий: {top_category}", font=("Arial", 14), text_color="#4CAF50").pack(anchor="w", pady=5)
        
        ctk.CTkLabel(scroll, text="✨ Рекомендуемые экскурсии", font=("Arial", 16, "bold")).pack(anchor="w", pady=(20, 10))
        
        if categories_interest:
            top_categories = list(categories_interest.keys())[:3]
            placeholders = ','.join(['%s'] * len(top_categories))
            
            recs = DB.query(
                f"""SELECT e.*, u.first_name, u.last_name, g.rating
                   FROM excursions e
                   JOIN users u ON e.guide_id = u.id
                   JOIN guides g ON u.id = g.user_id
                   WHERE e.status = 'PUB' AND e.id NOT IN (
                       SELECT s.excursion_id FROM bookings b
                       JOIN schedules s ON b.schedule_id = s.id
                       WHERE b.tourist_id = %s AND b.status = 'CONF'
                   )
                   AND e.category IN ({placeholders})
                   ORDER BY g.rating DESC, e.id DESC
                   LIMIT 6""",
                [user_id] + top_categories, fetch='all'
            )
        else:
            recs = DB.query(
                """SELECT e.*, u.first_name, u.last_name, g.rating,
                          (SELECT COUNT(*) FROM bookings b JOIN schedules s ON b.schedule_id = s.id WHERE s.excursion_id = e.id) as booking_count
                   FROM excursions e
                   JOIN users u ON e.guide_id = u.id
                   JOIN guides g ON u.id = g.user_id
                   WHERE e.status = 'PUB'
                   ORDER BY booking_count DESC, g.rating DESC
                   LIMIT 6""",
                fetch='all'
            )
        
        if recs:
            for r in recs:
                card = ctk.CTkFrame(scroll, fg_color="white", corner_radius=12)
                card.pack(fill="x", pady=8)
                
                ctk.CTkLabel(card, text=r['title'], font=("Arial", 15, "bold")).pack(anchor="w", padx=15, pady=(10, 5))
                ctk.CTkLabel(card, text=f"Гид: {r['first_name']} {r['last_name']} • {r['category']} • ★ {r['rating']}", 
                            font=("Arial", 11)).pack(anchor="w", padx=15)
                ctk.CTkLabel(card, text=r['description'][:100] + "..." if r['description'] and len(r['description']) > 100 else r['description'], 
                            wraplength=600, justify="left", font=("Arial", 11), text_color="#666").pack(anchor="w", padx=15, pady=5)
                
                btn_frame = ctk.CTkFrame(card, fg_color="transparent")
                btn_frame.pack(anchor="e", padx=15, pady=(0, 10))
                ctk.CTkButton(btn_frame, text="Забронировать", width=140, height=35, fg_color=BG_CARD, 
                             text_color="black", corner_radius=6,
                             command=lambda eid=r['id']: self._book_excursion(eid)).pack()
        else:
            ctk.CTkLabel(scroll, text="Пока нет рекомендаций. Посмотрите другие экскурсии!", font=("Arial", 12), text_color="gray").pack(pady=20)
    
    def _book_excursion(self, exc_id):
        e = DB.query("SELECT * FROM excursions WHERE id=%s", (exc_id,), fetch='one')
        if not e:
            messagebox.showerror("Ошибка", "Экскурсия не найдена")
            return
        
        self.master.log_activity(self.master.current_user['id'], exc_id, 'VIEW')
        
        schedules = DB.query(
            "SELECT id, start_datetime FROM schedules WHERE excursion_id=%s AND start_datetime > NOW()",
            (exc_id,)
        )
        
        if not schedules:
            messagebox.showerror("Ошибка", "Нет доступных дат")
            return
        
        win = ctk.CTkToplevel(self.master)
        win.title("Бронирование")
        win.geometry("420x420")
        win.configure(fg_color=BG_WHITE)
        win.transient(self.master)
        win.grab_set()
        
        ctk.CTkLabel(win, text=e['title'], font=("Arial", 18, "bold")).pack(pady=20)
        ctk.CTkLabel(win, text="Выберите дату:", font=("Arial", 13)).pack()
        
        schedule_var = ctk.StringVar(value=schedules[0]['start_datetime'].strftime('%d.%m.%Y %H:%M'))
        schedule_combo = ctk.CTkComboBox(win, variable=schedule_var,
                                        values=[s['start_datetime'].strftime('%d.%m.%Y %H:%M') for s in schedules],
                                        fg_color=BG_MAIN, width=280, height=35, corner_radius=8)
        schedule_combo.pack(pady=10)
        
        ctk.CTkLabel(win, text="Количество участников:", font=("Arial", 13)).pack()
        parts_var = ctk.IntVar(value=1)
        parts_spin = ctk.CTkEntry(win, textvariable=parts_var, width=100, height=35, fg_color=BG_MAIN, corner_radius=8)
        parts_spin.pack(pady=5)
        
        def confirm_booking():
            schedule_id = None
            for s in schedules:
                if s['start_datetime'].strftime('%d.%m.%Y %H:%M') == schedule_var.get():
                    schedule_id = s['id']
                    break
            if schedule_id:
                DB.query(
                    "INSERT INTO bookings (tourist_id, schedule_id, participants_count, status)"
                    " VALUES (%s,%s,%s,'PEND')",
                    (self.master.current_user['id'], schedule_id, parts_var.get()), fetch=None
                )
                self.master.log_activity(self.master.current_user['id'], exc_id, 'BOOK')
                self.master.add_notification(self.master.current_user['id'], 'Бронирование создано',
                                            f'Вы успешно забронировали экскурсию "{e["title"]}"', 'BOOKING')
                # Уведомление гиду о новом бронировании
                user = self.master.current_user
                self.master.add_notification(
                    e['guide_id'],
                    'Новое бронирование',
                    f'{user["first_name"]} {user["last_name"]} забронировал вашу экскурсию "{e["title"]}"',
                    'BOOKING'
                )
                messagebox.showinfo("Успех", "Бронирование создано!")
                win.destroy()
            else:
                messagebox.showerror("Ошибка", "Выберите дату")
        
        ctk.CTkButton(win, text="Подтвердить", width=200, height=40, fg_color=BG_CARD, text_color="black", corner_radius=8,
                     command=confirm_booking).pack(pady=25)


# Страница уведомлений
class NotificationsPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        top_frame = ctk.CTkFrame(self, fg_color="transparent")
        top_frame.pack(fill="x", pady=10, padx=5)
        
        ctk.CTkButton(top_frame, text="← Назад", fg_color=BG_CARD, text_color="black", 
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(side="left")
        
        ctk.CTkButton(top_frame, text="Отметить все прочитанными", fg_color=BG_CARD, text_color="black", 
                      width=200, height=35, corner_radius=8, command=self._mark_all_read).pack(side="right")
        
        ctk.CTkLabel(self, text="Уведомления", font=("Arial", 22, "bold")).pack(pady=15)
        
        self.scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.scroll.pack(fill="both", expand=True, padx=20, pady=10)
        
        self._load_notifications()
    
    def _load_notifications(self):
        for w in self.scroll.winfo_children():
            w.destroy()
        
        notifications = DB.query(
            "SELECT * FROM notifications WHERE user_id = %s ORDER BY created_at DESC",
            (self.master.current_user['id'],), fetch='all'
        )
        
        if not notifications:
            ctk.CTkLabel(self.scroll, text="У вас пока нет уведомлений", font=("Arial", 14), text_color="gray").pack(pady=30)
            return
        
        type_colors = {
            'INFO': "#2196F3",
            'BOOKING': "#4CAF50",
            'REMINDER': "#FF9800",
            'PROMO': "#9C27B0"
        }
        
        type_icons = {
            'INFO': "ℹ️",
            'BOOKING': "📅",
            'REMINDER': "⏰",
            'PROMO': "🎉"
        }
        
        for n in notifications:
            card = ctk.CTkFrame(self.scroll, fg_color="white", corner_radius=10)
            card.pack(fill="x", pady=6)

            if not n['is_read']:
                card.configure(border_width=2, border_color=type_colors.get(n['type'], "#666"))

            header = ctk.CTkFrame(card, fg_color="transparent")
            header.pack(fill="x", padx=15, pady=(10, 5))

            ctk.CTkLabel(header, text=f"{type_icons.get(n['type'], '📢')} {n['title']}",
                         font=("Arial", 14, "bold")).pack(side="left")
            ctk.CTkLabel(header, text=n['created_at'].strftime('%d.%m.%Y %H:%M'),
                         font=("Arial", 10), text_color="#999").pack(side="right")

            ctk.CTkLabel(card, text=n['message'], wraplength=700, justify="left",
                         font=("Arial", 12)).pack(anchor="w", padx=15, pady=(0, 8))

            btn_row = ctk.CTkFrame(card, fg_color="transparent")
            btn_row.pack(fill="x", padx=15, pady=(0, 10))

            ctk.CTkButton(btn_row, text="🗑️ Удалить", width=110, height=30, fg_color="#E57373",
                         text_color="white", corner_radius=6, hover_color="#C62828",
                         command=lambda nid=n['id']: self._delete_notification(nid)).pack(side="right", padx=(5, 0))

            if not n['is_read']:
                ctk.CTkButton(btn_row, text="Прочитано", width=120, height=30, fg_color=BG_CARD,
                             text_color="black", corner_radius=6,
                             command=lambda nid=n['id']: self._mark_read(nid)).pack(side="right")

    def _delete_notification(self, notif_id):
        DB.query("DELETE FROM notifications WHERE id=%s", (notif_id,), fetch=None)
        self._load_notifications()

    def _mark_read(self, notif_id):
        DB.query("UPDATE notifications SET is_read = TRUE WHERE id = %s", (notif_id,), fetch=None)
        self._load_notifications()

    def _mark_all_read(self):
        DB.query("UPDATE notifications SET is_read = TRUE WHERE user_id = %s", (self.master.current_user['id'],), fetch=None)
        self._load_notifications()
        messagebox.showinfo("Успех", "Все уведомления отмечены прочитанными")


# Страница отчетов (для админа)
class ReportsPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black", 
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        ctk.CTkLabel(self, text="Отчеты и аналитика", font=("Arial", 22, "bold")).pack(pady=15)
        
        self.tabview = ctk.CTkTabview(self, fg_color="white", corner_radius=12)
        self.tabview.pack(fill="both", expand=True, padx=20, pady=10)
        
        self.tabview.add("Общая статистика")
        self.tabview.add("Активность пользователей")
        self.tabview.add("Персональный отчет")
        
        self._load_general_stats()
        self._load_user_activity()
        self._load_personal_report()
    
    def _load_general_stats(self):
        tab = self.tabview.tab("Общая статистика")
        
        stats_frame = ctk.CTkFrame(tab, fg_color=BG_MAIN, corner_radius=12)
        stats_frame.pack(fill="x", padx=15, pady=15)
        
        total_users = DB.query("SELECT COUNT(*) as c FROM users", fetch='one')['c']
        total_guides = DB.query("SELECT COUNT(*) as c FROM guides", fetch='one')['c']
        total_excursions = DB.query("SELECT COUNT(*) as c FROM excursions WHERE status='PUB'", fetch='one')['c']
        total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings", fetch='one')['c']
        total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews", fetch='one')['c']
        total_activities = DB.query("SELECT COUNT(*) as c FROM user_activity", fetch='one')['c']
        
        stats = [
            ("👤 Всего пользователей", total_users),
            ("👤 Гидов", total_guides),
            ("🗺️ Экскурсий", total_excursions),
            ("📅 Бронирований", total_bookings),
            ("⭐ Отзывов", total_reviews),
            ("📊 Действий", total_activities),
        ]
        
        for i, (label, value) in enumerate(stats):
            box = ctk.CTkFrame(stats_frame, fg_color="white", corner_radius=10)
            box.grid(row=i//3, column=i%3, padx=10, pady=10, sticky="nsew")
            stats_frame.grid_columnconfigure(i%3, weight=1)
            ctk.CTkLabel(box, text=label, font=("Arial", 12), text_color="#666").pack(pady=(15, 5))
            ctk.CTkLabel(box, text=str(value), font=("Arial", 28, "bold")).pack(pady=(0, 15))
        
        export_frame = ctk.CTkFrame(tab, fg_color="transparent")
        export_frame.pack(pady=15)
        
        ctk.CTkButton(export_frame, text="📊 Экспорт статистики в PDF", fg_color=BG_CARD, text_color="black",
                     width=250, height=40, corner_radius=8, command=self._export_general_pdf).pack(side="left", padx=10)
        ctk.CTkButton(export_frame, text="📈 Экспорт статистики в Excel", fg_color=BG_CARD, text_color="black",
                     width=250, height=40, corner_radius=8, command=self._export_general_excel).pack(side="left", padx=10)
    
    def _load_user_activity(self):
        tab = self.tabview.tab("Активность пользователей")
        
        period_frame = ctk.CTkFrame(tab, fg_color=BG_MAIN, corner_radius=12)
        period_frame.pack(fill="x", padx=15, pady=15)
        
        ctk.CTkLabel(period_frame, text="Период:").pack(side="left", padx=15)
        
        self.activity_period = ctk.StringVar(value="30")
        for days, label in [("7", "7 дней"), ("30", "30 дней"), ("90", "90 дней"), ("365", "Год")]:
            ctk.CTkRadioButton(period_frame, text=label, variable=self.activity_period, value=days,
                              command=self._refresh_activity).pack(side="left", padx=10)
        
        self.activity_scroll = ctk.CTkScrollableFrame(tab, fg_color="transparent")
        self.activity_scroll.pack(fill="both", expand=True, padx=15, pady=10)
        
        export_btn = ctk.CTkButton(tab, text="📊 Экспорт активности в Excel", fg_color=BG_CARD, text_color="black",
                                   width=250, height=40, corner_radius=8, command=self._export_activity_excel)
        export_btn.pack(pady=10)
        
        self._refresh_activity()
    
    def _refresh_activity(self):
        for w in self.activity_scroll.winfo_children():
            w.destroy()
        
        days = int(self.activity_period.get())
        start_date = datetime.now() - timedelta(days=days)
        
        activity_by_day = DB.query(
            """SELECT DATE(created_at) as date, 
                      COUNT(CASE WHEN action_type='VIEW' THEN 1 END) as views,
                      COUNT(CASE WHEN action_type='BOOK' THEN 1 END) as bookings,
                      COUNT(CASE WHEN action_type='REVIEW' THEN 1 END) as reviews,
                      COUNT(CASE WHEN action_type='CANCEL' THEN 1 END) as cancels
               FROM user_activity 
               WHERE created_at > %s
               GROUP BY DATE(created_at)
               ORDER BY date DESC""",
            (start_date,), fetch='all'
        )
        
        ctk.CTkLabel(self.activity_scroll, text=f"Активность за последние {days} дней", 
                     font=("Arial", 14, "bold")).pack(anchor="w", pady=10)
        
        total_views = sum(a['views'] for a in activity_by_day)
        total_bookings = sum(a['bookings'] for a in activity_by_day)
        total_reviews = sum(a['reviews'] for a in activity_by_day)
        
        summary_frame = ctk.CTkFrame(self.activity_scroll, fg_color="white", corner_radius=10)
        summary_frame.pack(fill="x", pady=10)
        
        ctk.CTkLabel(summary_frame, text="📊 Сводка за период", font=("Arial", 12, "bold")).pack(anchor="w", padx=15, pady=5)
        summary_grid = ctk.CTkFrame(summary_frame, fg_color="transparent")
        summary_grid.pack(fill="x", padx=15, pady=5)
        
        ctk.CTkLabel(summary_grid, text=f"👁 Просмотров: {total_views}", font=("Arial", 11)).pack(side="left", padx=10)
        ctk.CTkLabel(summary_grid, text=f"📅 Бронирований: {total_bookings}", font=("Arial", 11)).pack(side="left", padx=10)
        ctk.CTkLabel(summary_grid, text=f"⭐ Отзывов: {total_reviews}", font=("Arial", 11)).pack(side="left", padx=10)
        
        for act in activity_by_day:
            card = ctk.CTkFrame(self.activity_scroll, fg_color="white", corner_radius=8)
            card.pack(fill="x", pady=5)
            
            date_str = act['date'].strftime('%d.%m.%Y') if act['date'] else "—"
            ctk.CTkLabel(card, text=date_str, font=("Arial", 13, "bold"), width=120).pack(side="left", padx=15, pady=10)
            ctk.CTkLabel(card, text=f"👁 {act['views']}", font=("Arial", 12)).pack(side="left", padx=15)
            ctk.CTkLabel(card, text=f"📅 {act['bookings']}", font=("Arial", 12)).pack(side="left", padx=15)
            ctk.CTkLabel(card, text=f"★ {act['reviews']}", font=("Arial", 12)).pack(side="left", padx=15)
        
        ctk.CTkLabel(self.activity_scroll, text="\n🏆 Самые активные пользователи", 
                     font=("Arial", 14, "bold")).pack(anchor="w", pady=(20, 10))
        
        active_users = DB.query(
            """SELECT u.id, u.first_name, u.last_name, u.email, u.role, COUNT(ua.id) as activity_count
               FROM users u
               JOIN user_activity ua ON u.id = ua.user_id
               WHERE ua.created_at > %s AND u.role != 'ADM'
               GROUP BY u.id
               ORDER BY activity_count DESC
               LIMIT 10""",
            (start_date,), fetch='all'
        )
        
        for u in active_users:
            card = ctk.CTkFrame(self.activity_scroll, fg_color=BG_MAIN, corner_radius=8)
            card.pack(fill="x", pady=5)
            role_icon = "👤" if u['role'] == 'TUR' else "🎓"
            ctk.CTkLabel(card, text=f"{role_icon} {u['first_name']} {u['last_name']}", font=("Arial", 13, "bold")).pack(anchor="w", padx=15, pady=5)
            ctk.CTkLabel(card, text=f"Email: {u['email']} • Действий: {u['activity_count']}", font=("Arial", 11)).pack(anchor="w", padx=15, pady=(0, 5))
    
    def _load_personal_report(self):
        tab = self.tabview.tab("Персональный отчет")
        
        user_frame = ctk.CTkFrame(tab, fg_color=BG_MAIN, corner_radius=12)
        user_frame.pack(fill="x", padx=15, pady=15)
        
        ctk.CTkLabel(user_frame, text="Выберите пользователя:").pack(side="left", padx=15)
        
        users = DB.query("SELECT id, first_name, last_name, email, role FROM users WHERE role != 'ADM'", fetch='all')
        user_names = [f"{u['first_name']} {u['last_name']} ({u['email']}) - {u['role']}" for u in users]
        user_names.insert(0, "Выберите пользователя...")
        
        self.selected_user_var = ctk.StringVar(value="Выберите пользователя...")
        self.user_combo = ctk.CTkComboBox(user_frame, variable=self.selected_user_var, values=user_names,
                                         fg_color="white", width=350, height=35, corner_radius=8)
        self.user_combo.pack(side="left", padx=15)
        
        ctk.CTkButton(user_frame, text="Сформировать отчет", fg_color=BG_CARD, text_color="black",
                     width=180, height=35, corner_radius=8, command=self._show_personal_report).pack(side="left", padx=15)
        
        self.personal_report_frame = ctk.CTkFrame(tab, fg_color="white", corner_radius=12)
        self.personal_report_frame.pack(fill="both", expand=True, padx=15, pady=15)
        
        ctk.CTkLabel(self.personal_report_frame, text="Выберите пользователя и нажмите 'Сформировать отчет'", 
                     font=("Arial", 12), text_color="gray").pack(pady=50)
    
    def _show_personal_report(self):
        for w in self.personal_report_frame.winfo_children():
            w.destroy()
        
        selected = self.selected_user_var.get()
        if selected == "Выберите пользователя...":
            ctk.CTkLabel(self.personal_report_frame, text="Пожалуйста, выберите пользователя", 
                         font=("Arial", 12), text_color="red").pack(pady=50)
            return
        
        users = DB.query("SELECT id, first_name, last_name, email, role FROM users WHERE role != 'ADM'", fetch='all')
        user_id = None
        user_info = None
        for u in users:
            if selected.startswith(f"{u['first_name']} {u['last_name']}"):
                user_id = u['id']
                user_info = u
                break
        
        if not user_id:
            return
        
        total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings WHERE tourist_id = %s", (user_id,), fetch='one')['c']
        total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews WHERE user_id = %s", (user_id,), fetch='one')['c']
        total_views = DB.query("SELECT COUNT(*) as c FROM user_activity WHERE user_id = %s AND action_type = 'VIEW'", (user_id,), fetch='one')['c']
        
        categories = DB.query(
            """SELECT e.category, COUNT(*) as count
               FROM bookings b
               JOIN schedules s ON b.schedule_id = s.id
               JOIN excursions e ON s.excursion_id = e.id
               WHERE b.tourist_id = %s AND b.status = 'CONF'
               GROUP BY e.category
               ORDER BY count DESC""",
            (user_id,), fetch='all'
        )
        
        activity_history = DB.query(
            """SELECT DATE(created_at) as date, COUNT(*) as count
               FROM user_activity
               WHERE user_id = %s
               GROUP BY DATE(created_at)
               ORDER BY date DESC
               LIMIT 30""",
            (user_id,), fetch='all'
        )
        
        scroll = ctk.CTkScrollableFrame(self.personal_report_frame, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=10, pady=10)
        
        info_card = ctk.CTkFrame(scroll, fg_color=BG_MAIN, corner_radius=12)
        info_card.pack(fill="x", pady=10)
        
        role_name = "Турист" if user_info['role'] == 'TUR' else "Гид"
        ctk.CTkLabel(info_card, text=f"📊 Отчет по пользователю: {user_info['first_name']} {user_info['last_name']} ({role_name})", 
                     font=("Arial", 16, "bold")).pack(anchor="w", padx=15, pady=10)
        ctk.CTkLabel(info_card, text=f"Email: {user_info['email']}", font=("Arial", 12)).pack(anchor="w", padx=15)
        ctk.CTkLabel(info_card, text=f"Дата регистрации: {user_info.get('registered_at', '—')}", font=("Arial", 12)).pack(anchor="w", padx=15, pady=(0, 10))
        
        stats_card = ctk.CTkFrame(scroll, fg_color=BG_MAIN, corner_radius=12)
        stats_card.pack(fill="x", pady=10)
        
        ctk.CTkLabel(stats_card, text="📈 Статистика активности", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=10)
        
        stats_grid = ctk.CTkFrame(stats_card, fg_color="transparent")
        stats_grid.pack(fill="x", padx=15, pady=10)
        
        stats_data = [
            ("Бронирований", total_bookings),
            ("Отзывов", total_reviews),
            ("Просмотров", total_views),
            ("Всего действий", total_bookings + total_reviews + total_views),
        ]
        
        for i, (label, value) in enumerate(stats_data):
            box = ctk.CTkFrame(stats_grid, fg_color="white", corner_radius=8)
            box.grid(row=0, column=i, padx=10, pady=5, sticky="nsew")
            stats_grid.grid_columnconfigure(i, weight=1)
            ctk.CTkLabel(box, text=label, font=("Arial", 11), text_color="#666").pack(pady=(10, 5))
            ctk.CTkLabel(box, text=str(value), font=("Arial", 20, "bold")).pack(pady=(0, 10))
        
        if categories:
            pref_card = ctk.CTkFrame(scroll, fg_color=BG_MAIN, corner_radius=12)
            pref_card.pack(fill="x", pady=10)
            
            ctk.CTkLabel(pref_card, text="🎯 Предпочтения по категориям", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=10)
            
            for cat in categories:
                bar_frame = ctk.CTkFrame(pref_card, fg_color="transparent")
                bar_frame.pack(fill="x", padx=15, pady=5)
                
                ctk.CTkLabel(bar_frame, text=cat['category'], width=150, anchor="w").pack(side="left")
                
                bar_bg = ctk.CTkFrame(bar_frame, fg_color=BG_WHITE, height=20, width=300, corner_radius=4)
                bar_bg.pack(side="left", padx=10)
                
                max_count = max([c['count'] for c in categories])
                width_percent = (cat['count'] / max_count) * 300 if max_count > 0 else 0
                
                bar_fill = ctk.CTkFrame(bar_bg, fg_color=COLOR_ACCENT, height=20, width=width_percent, corner_radius=4)
                bar_fill.place(x=0, y=0)
                
                ctk.CTkLabel(bar_frame, text=f"{cat['count']} экскурсий", width=100).pack(side="left")
            
            top_category = categories[0]['category']
            ctk.CTkLabel(pref_card, text=f"\n💡 Рекомендация: Пользователь проявляет наибольший интерес к категории '{top_category}'", 
                         font=("Arial", 12), text_color="#4CAF50").pack(anchor="w", padx=15, pady=(10, 15))
        
        if activity_history:
            activity_card = ctk.CTkFrame(scroll, fg_color=BG_MAIN, corner_radius=12)
            activity_card.pack(fill="x", pady=10)
            
            ctk.CTkLabel(activity_card, text="📅 График активности (по дням)", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=10)
            
            hist_frame = ctk.CTkFrame(activity_card, fg_color="transparent")
            hist_frame.pack(fill="x", padx=15, pady=10)
            
            max_count = max([a['count'] for a in activity_history]) if activity_history else 1
            for a in activity_history[:14]:
                bar_frame = ctk.CTkFrame(hist_frame, fg_color="transparent")
                bar_frame.pack(fill="x", pady=3)
                
                date_str = a['date'].strftime('%d.%m') if a['date'] else "—"
                ctk.CTkLabel(bar_frame, text=date_str, width=80, anchor="w").pack(side="left")
                
                bar_bg = ctk.CTkFrame(bar_frame, fg_color=BG_WHITE, height=15, width=300, corner_radius=3)
                bar_bg.pack(side="left", padx=10)
                
                width_percent = (a['count'] / max_count) * 300 if max_count > 0 else 0
                bar_fill = ctk.CTkFrame(bar_bg, fg_color=COLOR_ACCENT, height=15, width=width_percent, corner_radius=3)
                bar_fill.place(x=0, y=0)
                
                ctk.CTkLabel(bar_frame, text=str(a['count']), width=50).pack(side="left")
        
        export_frame = ctk.CTkFrame(scroll, fg_color="transparent")
        export_frame.pack(pady=15)
        
        ctk.CTkButton(export_frame, text="📊 Экспорт персонального отчета в PDF", fg_color=BG_CARD, text_color="black",
                     width=250, height=40, corner_radius=8, 
                     command=lambda: self._export_personal_pdf(user_id, user_info)).pack(side="left", padx=10)
        ctk.CTkButton(export_frame, text="📈 Экспорт персонального отчета в Excel", fg_color=BG_CARD, text_color="black",
                     width=250, height=40, corner_radius=8,
                     command=lambda: self._export_personal_excel(user_id, user_info)).pack(side="left", padx=10)
    
    # Методы экспорта
    
    def _get_pdf_styles(self):
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CustomTitle', 
            parent=styles['Heading1'], 
            fontSize=14, 
            alignment=TA_CENTER,
            fontName=FONT_BOLD
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=10,
            alignment=TA_LEFT,
            fontName=FONT_BOLD
        )
        
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=8,
            fontName=FONT_NAME
        )
        
        return styles, title_style, heading_style, normal_style
    
    def _export_general_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Ошибка", "Для экспорта в PDF установите reportlab: pip install reportlab")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")],
            initialfile="statistics_report.pdf"
        )
        
        if not file_path:
            return
        
        try:
            doc = SimpleDocTemplate(file_path, pagesize=A4)
            styles, title_style, heading_style, normal_style = self._get_pdf_styles()
            story = []
            
            story.append(Paragraph("Отчет по статистике Guide-Advisor", title_style))
            story.append(Spacer(1, 0.2*inch))
            story.append(Paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}", normal_style))
            story.append(Spacer(1, 0.1*inch))
            
            total_users = DB.query("SELECT COUNT(*) as c FROM users", fetch='one')['c']
            total_guides = DB.query("SELECT COUNT(*) as c FROM guides", fetch='one')['c']
            total_excursions = DB.query("SELECT COUNT(*) as c FROM excursions WHERE status='PUB'", fetch='one')['c']
            total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings", fetch='one')['c']
            total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews", fetch='one')['c']
            total_activities = DB.query("SELECT COUNT(*) as c FROM user_activity", fetch='one')['c']
            
            story.append(Paragraph("Общая статистика системы", heading_style))
            story.append(Spacer(1, 0.05*inch))
            
            data = [
                ["Показатель", "Значение"],
                ["Всего пользователей", str(total_users)],
                ["Гидов", str(total_guides)],
                ["Экскурсий", str(total_excursions)],
                ["Бронирований", str(total_bookings)],
                ["Отзывов", str(total_reviews)],
                ["Действий пользователей", str(total_activities)],
            ]
            
            table = Table(data, colWidths=[2.5*inch, 1.2*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#4A90D9')),
                ('TEXTCOLOR', (0, 0), (1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (1, 0), FONT_BOLD),
                ('FONTSIZE', (0, 0), (1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (1, 0), 8),
                ('BACKGROUND', (0, 1), (1, -1), colors.HexColor('#F5F5F5')),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#Cccccc')),
                ('FONTNAME', (0, 1), (1, -1), FONT_NAME),
                ('FONTSIZE', (0, 1), (1, -1), 8),
            ]))
            story.append(table)
            
            story.append(PageBreak())
            
            # Топ 10 экскурсий - увеличенный столбец "бронирований"
            story.append(Paragraph("ТОП-10 популярных экскурсий", heading_style))
            story.append(Spacer(1, 0.05*inch))
            
            pop_exc = DB.query(
                """SELECT e.title, e.category, COUNT(DISTINCT b.id) as booking_count,
                          COUNT(DISTINCT r.id) as review_count, AVG(r.rating) as avg_rating
                   FROM excursions e
                   LEFT JOIN schedules s ON e.id = s.excursion_id
                   LEFT JOIN bookings b ON s.id = b.schedule_id
                   LEFT JOIN reviews r ON e.id = r.excursion_id
                   WHERE e.status = 'PUB'
                   GROUP BY e.id
                   ORDER BY booking_count DESC
                   LIMIT 10""",
                fetch='all'
            )
            
            # Увеличенная ширина столбца "Бронирований"
            exc_data = [["Название", "Категория", "Бронирований", "Отзывов", "Рейтинг"]]
            for e in pop_exc:
                avg = f"{round(e['avg_rating'], 1)}" if e['avg_rating'] else "Нет"
                exc_data.append([e['title'][:35], e['category'], str(e['booking_count']), str(e['review_count']), avg])
            
            exc_table = Table(exc_data, colWidths=[2.2*inch, 1*inch, 1.0*inch, 0.8*inch, 0.7*inch])  # Увеличен столбец "Бронирований"
            exc_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4A90D9')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), FONT_BOLD),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#Cccccc')),
                ('FONTNAME', (0, 1), (-1, -1), FONT_NAME),
            ]))
            story.append(exc_table)
            
            doc.build(story)
            messagebox.showinfo("Успех", f"PDF отчет сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать PDF: {e}")
    
    def _export_general_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Ошибка", "Для экспорта в Excel установите openpyxl: pip install openpyxl")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="statistics_report.xlsx"
        )
        
        if not file_path:
            return
        
        try:
            wb = openpyxl.Workbook()
            
            ws1 = wb.active
            ws1.title = "Общая статистика"
            
            total_users = DB.query("SELECT COUNT(*) as c FROM users", fetch='one')['c']
            total_guides = DB.query("SELECT COUNT(*) as c FROM guides", fetch='one')['c']
            total_excursions = DB.query("SELECT COUNT(*) as c FROM excursions WHERE status='PUB'", fetch='one')['c']
            total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings", fetch='one')['c']
            total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews", fetch='one')['c']
            
            headers = ["Показатель", "Значение"]
            ws1.append(headers)
            ws1.append(["Всего пользователей", total_users])
            ws1.append(["Гидов", total_guides])
            ws1.append(["Экскурсий", total_excursions])
            ws1.append(["Бронирований", total_bookings])
            ws1.append(["Отзывов", total_reviews])
            
            header_font = Font(bold=True, size=12, color="FFFFFF")
            header_fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
            for cell in ws1[1]:
                cell.font = header_font
                cell.fill = header_fill
            
            ws2 = wb.create_sheet("ТОП-10 экскурсий")
            excursions_stats = DB.query(
                """SELECT e.title, e.category, COUNT(DISTINCT b.id) as booking_count,
                          COUNT(DISTINCT r.id) as review_count, AVG(r.rating) as avg_rating
                   FROM excursions e
                   LEFT JOIN schedules s ON e.id = s.excursion_id
                   LEFT JOIN bookings b ON s.id = b.schedule_id
                   LEFT JOIN reviews r ON e.id = r.excursion_id
                   WHERE e.status = 'PUB'
                   GROUP BY e.id
                   ORDER BY booking_count DESC
                   LIMIT 10""",
                fetch='all'
            )
            
            exc_headers = ["Название", "Категория", "Бронирований", "Отзывов", "Средний рейтинг"]
            ws2.append(exc_headers)
            for e in excursions_stats:
                avg = round(e['avg_rating'], 1) if e['avg_rating'] else "Нет"
                ws2.append([e['title'], e['category'], e['booking_count'], e['review_count'], avg])
            
            for cell in ws2[1]:
                cell.font = header_font
                cell.fill = header_fill
            
            # Увеличенная ширина столбца "Бронирований"
            ws2.column_dimensions['A'].width = 45
            ws2.column_dimensions['B'].width = 20
            ws2.column_dimensions['C'].width = 18  # Увеличен столбец "Бронирований"
            ws2.column_dimensions['D'].width = 15
            ws2.column_dimensions['E'].width = 18
            
            for ws in [ws1]:
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(file_path)
            messagebox.showinfo("Успех", f"Excel отчет сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel: {e}")
    
    def _export_activity_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Ошибка", "Для экспорта в Excel установите openpyxl")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="user_activity_report.xlsx"
        )
        
        if not file_path:
            return
        
        try:
            wb = openpyxl.Workbook()
            
            days = int(self.activity_period.get())
            start_date = datetime.now() - timedelta(days=days)
            
            ws1 = wb.active
            ws1.title = "Активность по дням"
            
            activity_by_day = DB.query(
                """SELECT DATE(created_at) as date, 
                          COUNT(CASE WHEN action_type='VIEW' THEN 1 END) as views,
                          COUNT(CASE WHEN action_type='BOOK' THEN 1 END) as bookings,
                          COUNT(CASE WHEN action_type='REVIEW' THEN 1 END) as reviews,
                          COUNT(CASE WHEN action_type='CANCEL' THEN 1 END) as cancels
                   FROM user_activity 
                   WHERE created_at > %s
                   GROUP BY DATE(created_at)
                   ORDER BY date DESC""",
                (start_date,), fetch='all'
            )
            
            headers = ["Дата", "Просмотры", "Бронирования", "Отзывы", "Отмены"]
            ws1.append(headers)
            for act in activity_by_day:
                date_str = act['date'].strftime('%d.%m.%Y') if act['date'] else "—"
                ws1.append([date_str, act['views'], act['bookings'], act['reviews'], act['cancels']])
            
            ws2 = wb.create_sheet("Активные пользователи")
            
            active_users = DB.query(
                """SELECT u.id, u.first_name, u.last_name, u.email, u.role, COUNT(ua.id) as activity_count
                   FROM users u
                   JOIN user_activity ua ON u.id = ua.user_id
                   WHERE ua.created_at > %s AND u.role != 'ADM'
                   GROUP BY u.id
                   ORDER BY activity_count DESC""",
                (start_date,), fetch='all'
            )
            
            ws2.append(["Имя", "Фамилия", "Email", "Роль", "Количество действий"])
            for u in active_users:
                role_rus = "Турист" if u['role'] == 'TUR' else "Гид"
                ws2.append([u['first_name'], u['last_name'], u['email'], role_rus, u['activity_count']])
            
            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
            for ws in [ws1, ws2]:
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 30)
                    ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(file_path)
            messagebox.showinfo("Успех", f"Excel отчет сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel: {e}")
    
    def _export_personal_pdf(self, user_id, user_info):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Ошибка", "Для экспорта в PDF установите reportlab")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")],
            initialfile=f"personal_report_{user_info['first_name']}_{user_info['last_name']}.pdf"
        )
        
        if not file_path:
            return
        
        try:
            doc = SimpleDocTemplate(file_path, pagesize=A4)
            styles, title_style, heading_style, normal_style = self._get_pdf_styles()
            story = []
            
            role_name = "Турист" if user_info['role'] == 'TUR' else "Гид"
            story.append(Paragraph(f"Персональный отчет: {user_info['first_name']} {user_info['last_name']} ({role_name})", title_style))
            story.append(Spacer(1, 0.2*inch))
            story.append(Paragraph(f"Email: {user_info['email']}", normal_style))
            story.append(Paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}", normal_style))
            story.append(Spacer(1, 0.1*inch))
            
            total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings WHERE tourist_id = %s", (user_id,), fetch='one')['c']
            total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews WHERE user_id = %s", (user_id,), fetch='one')['c']
            total_views = DB.query("SELECT COUNT(*) as c FROM user_activity WHERE user_id = %s AND action_type = 'VIEW'", (user_id,), fetch='one')['c']
            
            story.append(Paragraph("Статистика активности", heading_style))
            story.append(Spacer(1, 0.05*inch))
            
            data = [
                ["Показатель", "Значение"],
                ["Бронирований", str(total_bookings)],
                ["Отзывов", str(total_reviews)],
                ["Просмотров", str(total_views)],
                ["Всего действий", str(total_bookings + total_reviews + total_views)],
            ]
            
            table = Table(data, colWidths=[2*inch, 1.2*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#4A90D9')),
                ('TEXTCOLOR', (0, 0), (1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (1, 0), FONT_BOLD),
                ('FONTSIZE', (0, 0), (1, 0), 9),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#Cccccc')),
                ('FONTNAME', (0, 1), (1, -1), FONT_NAME),
                ('FONTSIZE', (0, 1), (1, -1), 7),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.1*inch))
            
            categories = DB.query(
                """SELECT e.category, COUNT(*) as count
                   FROM bookings b
                   JOIN schedules s ON b.schedule_id = s.id
                   JOIN excursions e ON s.excursion_id = e.id
                   WHERE b.tourist_id = %s AND b.status = 'CONF'
                   GROUP BY e.category
                   ORDER BY count DESC""",
                (user_id,), fetch='all'
            )
            
            if categories:
                story.append(Paragraph("Предпочтения по категориям", heading_style))
                cat_data = [["Категория", "Количество экскурсий"]]
                for cat in categories:
                    cat_data.append([cat['category'], str(cat['count'])])
                
                cat_table = Table(cat_data, colWidths=[2*inch, 1.2*inch])
                cat_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#4A90D9')),
                    ('TEXTCOLOR', (0, 0), (1, 0), colors.whitesmoke),
                    ('FONTNAME', (0, 0), (1, 0), FONT_BOLD),
                    ('FONTSIZE', (0, 0), (1, 0), 9),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#Cccccc')),
                    ('FONTNAME', (0, 1), (1, -1), FONT_NAME),
                    ('FONTSIZE', (0, 1), (1, -1), 7),
                ]))
                story.append(cat_table)
                
                top_category = categories[0]['category']
                story.append(Spacer(1, 0.05*inch))
                story.append(Paragraph(f"Рекомендация: Наибольший интерес к категории '{top_category}'", normal_style))
            
            doc.build(story)
            messagebox.showinfo("Успех", f"PDF отчет сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать PDF: {e}")
    
    def _export_personal_excel(self, user_id, user_info):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Ошибка", "Для экспорта в Excel установите openpyxl")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"personal_report_{user_info['first_name']}_{user_info['last_name']}.xlsx"
        )
        
        if not file_path:
            return
        
        try:
            wb = openpyxl.Workbook()
            
            ws1 = wb.active
            ws1.title = "Статистика"
            
            total_bookings = DB.query("SELECT COUNT(*) as c FROM bookings WHERE tourist_id = %s", (user_id,), fetch='one')['c']
            total_reviews = DB.query("SELECT COUNT(*) as c FROM reviews WHERE user_id = %s", (user_id,), fetch='one')['c']
            total_views = DB.query("SELECT COUNT(*) as c FROM user_activity WHERE user_id = %s AND action_type = 'VIEW'", (user_id,), fetch='one')['c']
            
            ws1.append(["Показатель", "Значение"])
            ws1.append(["Бронирований", total_bookings])
            ws1.append(["Отзывов", total_reviews])
            ws1.append(["Просмотров", total_views])
            ws1.append(["Всего действий", total_bookings + total_reviews + total_views])
            
            ws2 = wb.create_sheet("Категории")
            
            categories = DB.query(
                """SELECT e.category, COUNT(*) as count
                   FROM bookings b
                   JOIN schedules s ON b.schedule_id = s.id
                   JOIN excursions e ON s.excursion_id = e.id
                   WHERE b.tourist_id = %s AND b.status = 'CONF'
                   GROUP BY e.category
                   ORDER BY count DESC""",
                (user_id,), fetch='all'
            )
            
            ws2.append(["Категория", "Количество экскурсий"])
            for cat in categories:
                ws2.append([cat['category'], cat['count']])
            
            ws3 = wb.create_sheet("История бронирований")
            
            bookings = DB.query(
                """SELECT e.title, s.start_datetime, b.participants_count, b.status, b.booked_at
                   FROM bookings b
                   JOIN schedules s ON b.schedule_id = s.id
                   JOIN excursions e ON s.excursion_id = e.id
                   WHERE b.tourist_id = %s
                   ORDER BY s.start_datetime DESC""",
                (user_id,), fetch='all'
            )
            
            ws3.append(["Экскурсия", "Дата и время", "Участников", "Статус", "Дата бронирования"])
            for b in bookings:
                ws3.append([
                    b['title'],
                    b['start_datetime'].strftime('%d.%m.%Y %H:%M') if b['start_datetime'] else "—",
                    b['participants_count'],
                    b['status'],
                    b['booked_at'].strftime('%d.%m.%Y %H:%M') if b['booked_at'] else "—"
                ])
            
            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
            for ws in [ws1, ws2, ws3]:
                if ws.max_row >= 1:
                    for cell in ws[1]:
                        cell.font = header_font
                        cell.fill = header_fill
                    
                    for column in ws.columns:
                        max_length = 0
                        column_letter = column[0].column_letter
                        for cell in column:
                            try:
                                if len(str(cell.value)) > max_length:
                                    max_length = len(str(cell.value))
                            except:
                                pass
                        adjusted_width = min(max_length + 2, 40)
                        ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(file_path)
            messagebox.showinfo("Успех", f"Excel отчет сохранен: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать Excel: {e}")


# Каталог гидов
class GuidesCatalogPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← Панель управления", fg_color=BG_CARD, text_color="black", 
                      width=180, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        ctk.CTkLabel(self, text="Информация о гидах", font=("Arial", 22, "bold")).pack(pady=15)
        
        rows = DB.query(
            """SELECT u.id, u.first_name, u.last_name, u.username, g.specialization, g.rating, g.user_id, g.photo_path,
                      (SELECT COUNT(*) FROM excursions WHERE guide_id=g.user_id) as exc_count
               FROM guides g JOIN users u ON u.id=g.user_id ORDER BY g.rating DESC"""
        )
        
        canvas_frame = ctk.CTkScrollableFrame(self, fg_color="transparent")
        canvas_frame.pack(fill="both", expand=True, padx=15, pady=10)
        
        for r in rows:
            card = ctk.CTkFrame(canvas_frame, fg_color="white", corner_radius=12)
            card.pack(fill="x", pady=8)
            
            left = ctk.CTkFrame(card, fg_color="transparent")
            left.pack(side="left", padx=15, pady=10)
            
            photo_shown = False
            if r.get('photo_path') and os.path.exists(r['photo_path']):
                try:
                    from PIL import Image as PILImage, ImageDraw
                    img = PILImage.open(r['photo_path']).resize((160, 160)).convert("RGBA")
                    mask = PILImage.new("L", (160, 160), 0)
                    ImageDraw.Draw(mask).rounded_rectangle((0, 0, 160, 160), radius=20, fill=255)
                    img.putalpha(mask)
                    _img = ctk.CTkImage(img, size=(160, 160))
                    lbl = ctk.CTkLabel(left, image=_img, text="", width=160, height=160, fg_color="transparent")
                    lbl._img_ref = _img
                    lbl.pack()
                    photo_shown = True
                except Exception:
                    pass
            if not photo_shown:
                PhotoPlaceholder(left, width=160, height=160, text=f"{r['first_name'][0]}{r['last_name'][0]}")
            
            info = ctk.CTkFrame(card, fg_color="transparent")
            info.pack(side="left", fill="both", expand=True, padx=10, pady=10)
            
            ctk.CTkLabel(info, text=f"{r['first_name']} {r['last_name']}", font=("Arial", 16, "bold")).pack(anchor="w")
            ctk.CTkLabel(info, text=f"@{r['username']}", text_color="#666666").pack(anchor="w")
            ctk.CTkLabel(info, text=f"Специализация: {r['specialization'] or '—'}").pack(anchor="w")
            ctk.CTkLabel(info, text=f"★ {r['rating']:.1f}  •  {r['exc_count']} экскурсий").pack(anchor="w")
            
            btn_frame = ctk.CTkFrame(card, fg_color="transparent")
            btn_frame.pack(side="right", padx=15, pady=10)
            ctk.CTkButton(btn_frame, text="Профиль", width=120, height=35, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda gid=r['id']: self.master.show_guide_profile(gid)).pack()


# Каталог экскурсий
class ExcursionsCatalogPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← Панель управления", fg_color=BG_CARD, text_color="black", 
                      width=180, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        ctk.CTkLabel(self, text="Доступные экскурсии", font=("Arial", 22, "bold")).pack(pady=15)
        
        rows = DB.query(
            """SELECT e.id, e.title, e.category, e.duration_hours, e.description, e.max_participants,
                      e.photo_path, u.first_name, u.last_name
               FROM excursions e
               LEFT JOIN users u ON e.guide_id = u.id
               WHERE e.status = 'PUB'
               ORDER BY e.id DESC"""
        )

        canvas_frame = ctk.CTkScrollableFrame(self, fg_color="transparent")
        canvas_frame.pack(fill="both", expand=True, padx=15, pady=10)

        for r in rows:
            card = ctk.CTkFrame(canvas_frame, fg_color="white", corner_radius=12)
            card.pack(fill="x", pady=8)

            card_inner = ctk.CTkFrame(card, fg_color="transparent")
            card_inner.pack(fill="both", expand=True)

            photo_frame = ctk.CTkFrame(card_inner, fg_color="transparent", width=140)
            photo_frame.pack(side="left", fill="y", padx=15, pady=12)
            photo_frame.pack_propagate(False)

            photo_shown = False
            if r.get('photo_path') and os.path.exists(r['photo_path']):
                try:
                    from PIL import Image as PILImage, ImageDraw
                    pil_img = PILImage.open(r['photo_path'])
                    w, h = pil_img.size
                    side = min(w, h)
                    pil_img = pil_img.crop(((w - side) // 2, (h - side) // 2, (w + side) // 2, (h + side) // 2))
                    pil_img = pil_img.resize((130, 130)).convert("RGBA")
                    mask = PILImage.new("L", (130, 130), 0)
                    ImageDraw.Draw(mask).rounded_rectangle((0, 0, 130, 130), radius=14, fill=255)
                    pil_img.putalpha(mask)
                    _img = ctk.CTkImage(pil_img, size=(130, 130))
                    lbl = ctk.CTkLabel(photo_frame, image=_img, text="", width=130, height=130, fg_color="transparent")
                    lbl._img_ref = _img
                    lbl.pack(expand=True)
                    photo_shown = True
                except Exception:
                    pass
            if not photo_shown:
                ctk.CTkLabel(photo_frame, text="🖼️", width=130, height=130,
                             fg_color=COLOR_ACCENT, corner_radius=14, font=("Arial", 32)).pack(expand=True)

            info_frame = ctk.CTkFrame(card_inner, fg_color="transparent")
            info_frame.pack(side="left", fill="both", expand=True, pady=12)

            ctk.CTkLabel(info_frame, text=r['title'], font=("Arial", 16, "bold")).pack(anchor="w")
            ctk.CTkLabel(info_frame, text=f"Гид: {r['first_name']} {r['last_name']}  •  {r['category']}", font=("Arial", 12)).pack(anchor="w")
            ctk.CTkLabel(info_frame, text=f"Длительность: {r['duration_hours']} ч  •  Макс: {r['max_participants']} чел.", font=("Arial", 12)).pack(anchor="w")
            ctk.CTkLabel(info_frame, text=r['description'] or "Нет описания", wraplength=600, justify="left", font=("Arial", 11)).pack(anchor="w", pady=(4, 0))
            
            if self.master.current_user and self.master.current_user['role'] == 'TUR':
                btn_frame = ctk.CTkFrame(card, fg_color="transparent")
                btn_frame.pack(anchor="e", padx=15, pady=12)
                ctk.CTkButton(btn_frame, text="Забронировать", width=160, height=38, fg_color=BG_CARD, text_color="black", corner_radius=6,
                             command=lambda eid=r['id']: self._book_excursion(eid)).pack()
                ctk.CTkButton(btn_frame, text="⭐ Отзывы", width=140, height=38, fg_color=BG_MAIN, text_color="black", corner_radius=6,
                             command=lambda eid=r['id']: self.master.show_reviews(eid)).pack(pady=(5, 0))
    
    def _book_excursion(self, exc_id):
        e = DB.query("SELECT * FROM excursions WHERE id=%s", (exc_id,), fetch='one')
        if not e:
            messagebox.showerror("Ошибка", "Экскурсия не найдена")
            return
        
        self.master.log_activity(self.master.current_user['id'], exc_id, 'VIEW')
        
        schedules = DB.query(
            "SELECT id, start_datetime FROM schedules WHERE excursion_id=%s AND start_datetime > NOW()",
            (exc_id,)
        )
        
        if not schedules:
            messagebox.showerror("Ошибка", "Нет доступных дат")
            return
        
        win = ctk.CTkToplevel(self.master)
        win.title("Бронирование")
        win.geometry("420x420")
        win.configure(fg_color=BG_WHITE)
        win.transient(self.master)
        win.grab_set()
        
        ctk.CTkLabel(win, text=e['title'], font=("Arial", 18, "bold")).pack(pady=20)
        ctk.CTkLabel(win, text="Выберите дату:", font=("Arial", 13)).pack()
        
        schedule_var = ctk.StringVar(value=schedules[0]['start_datetime'].strftime('%d.%m.%Y %H:%M'))
        schedule_combo = ctk.CTkComboBox(win, variable=schedule_var,
                                        values=[s['start_datetime'].strftime('%d.%m.%Y %H:%M') for s in schedules],
                                        fg_color=BG_MAIN, width=280, height=35, corner_radius=8)
        schedule_combo.pack(pady=10)
        
        ctk.CTkLabel(win, text="Количество участников:", font=("Arial", 13)).pack()
        parts_var = ctk.IntVar(value=1)
        parts_spin = ctk.CTkEntry(win, textvariable=parts_var, width=100, height=35, fg_color=BG_MAIN, corner_radius=8)
        parts_spin.pack(pady=5)
        
        def confirm_booking():
            schedule_id = None
            for s in schedules:
                if s['start_datetime'].strftime('%d.%m.%Y %H:%M') == schedule_var.get():
                    schedule_id = s['id']
                    break
            if schedule_id:
                DB.query(
                    "INSERT INTO bookings (tourist_id, schedule_id, participants_count, status)"
                    " VALUES (%s,%s,%s,'PEND')",
                    (self.master.current_user['id'], schedule_id, parts_var.get()), fetch=None
                )
                self.master.log_activity(self.master.current_user['id'], exc_id, 'BOOK')
                self.master.add_notification(self.master.current_user['id'], 'Бронирование создано',
                                            f'Вы успешно забронировали экскурсию "{e["title"]}"', 'BOOKING')
                # Уведомление гиду о новом бронировании
                user = self.master.current_user
                self.master.add_notification(
                    e['guide_id'],
                    'Новое бронирование',
                    f'{user["first_name"]} {user["last_name"]} забронировал вашу экскурсию "{e["title"]}"',
                    'BOOKING'
                )
                messagebox.showinfo("Успех", "Бронирование создано!")
                win.destroy()
            else:
                messagebox.showerror("Ошибка", "Выберите дату")
        
        ctk.CTkButton(win, text="Подтвердить", width=200, height=40, fg_color=BG_CARD, text_color="black", corner_radius=8,
                     command=confirm_booking).pack(pady=25)


# Профиль гида
class GuideProfilePage(ctk.CTkFrame):
    def __init__(self, master, guide_id):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        self.guide_id = guide_id
        
        ctk.CTkButton(self, text="← В меню", fg_color=BG_CARD, text_color="black", 
                      width=150, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        guide = DB.query("SELECT * FROM guides WHERE user_id=%s", (guide_id,), fetch='one')
        user = DB.query("SELECT * FROM users WHERE id=%s", (guide_id,), fetch='one')
        
        if not guide or not user:
            ctk.CTkLabel(self, text="Гид не найден", text_color="red").pack(pady=20)
            return
        
        main_card = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        main_card.pack(fill="both", expand=True, padx=15, pady=15)
        
        ctk.CTkLabel(main_card, text=f"Профиль гида: {user['first_name']} {user['last_name']}", 
                     font=("Arial", 22, "bold")).pack(pady=20)
        
        cols = ctk.CTkFrame(main_card, fg_color="transparent")
        cols.pack(fill="both", expand=True, padx=25, pady=10)
        
        left = ctk.CTkFrame(cols, fg_color=BG_WHITE, corner_radius=12, width=250)
        left.pack(side="left", fill="y", padx=(0, 15), pady=10)

        photo_shown = False
        if guide.get('photo_path') and os.path.exists(guide['photo_path']):
            try:
                from PIL import Image as PILImage
                img = PILImage.open(guide['photo_path']).resize((200, 200))
                _img = ctk.CTkImage(img, size=(200, 200))
                img_lbl = ctk.CTkLabel(left, image=_img, text="", width=200, height=200)
                img_lbl._img_ref = _img
                img_lbl.pack(pady=25, padx=20)
                photo_shown = True
            except Exception:
                pass
        if not photo_shown:
            PhotoPlaceholder(left, width=200, height=200, text=f"{user['first_name'][0]}{user['last_name'][0]}").pack(pady=25, padx=20)
        
        right = ctk.CTkFrame(cols, fg_color=BG_WHITE, corner_radius=12)
        right.pack(side="left", expand=True, fill="both", pady=10)
        
        info = [
            f"ФИО: {user['first_name']} {user['last_name']}",
            f"Никнейм: @{user['username']}",
            f"Email: {user['email']}",
            f"Специализация: {guide['specialization'] or '—'}",
            f"Опыт: {guide['experience_years']} лет",
            f"Рейтинг: ★ {guide['rating']:.1f}",
            f"Описание: {guide['description'] or '—'}"
        ]
        for i in info:
            ctk.CTkLabel(right, text=i, font=("Arial", 13)).pack(anchor="w", padx=20, pady=6)
        
        ctk.CTkLabel(right, text="Экскурсии гида:", font=("Arial", 15, "bold")).pack(anchor="w", padx=20, pady=(20, 8))
        
        excursions = DB.query(
            "SELECT id, title, category, status FROM excursions WHERE guide_id=%s",
            (guide_id,)
        )
        for e in excursions:
            status_icon = "●" if e['status'] == 'PUB' else "●"
            ctk.CTkLabel(right, text=f"{status_icon} {e['title']} ({e['category']})", font=("Arial", 12)).pack(anchor="w", padx=35, pady=3)


# Профиль туриста
class TouristProfilePage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← В меню", fg_color=BG_CARD, text_color="black", 
                      width=150, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        
        u = self.master.current_user
        
        main_card = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        main_card.pack(fill="both", expand=True, padx=15, pady=15)
        
        ctk.CTkLabel(main_card, text=f"Профиль туриста: {u['first_name']} {u['last_name']}", 
                     font=("Arial", 22, "bold")).pack(pady=20)
        
        cols = ctk.CTkFrame(main_card, fg_color="transparent")
        cols.pack(fill="both", expand=True, padx=25, pady=10)
        
        left = ctk.CTkFrame(cols, fg_color=BG_WHITE, corner_radius=12, width=250)
        left.pack(side="left", fill="y", padx=(0, 15), pady=10)
        left.pack_propagate(False)

        self._tourist_photo_label = ctk.CTkLabel(left, text=f"{u['first_name'][0]}{u['last_name'][0]}",
                                                  width=200, height=200, fg_color=COLOR_ACCENT,
                                                  corner_radius=8, font=("Arial", 14))
        self._tourist_photo_label.pack(pady=(25, 10), padx=20)

        saved_photo = DB.query("SELECT photo_path FROM users WHERE id=%s", (u['id'],), fetch='one')
        if saved_photo and saved_photo.get('photo_path') and os.path.exists(saved_photo['photo_path']):
            self._load_tourist_photo(saved_photo['photo_path'])

        ctk.CTkButton(left, text="📷 Выбрать фото", fg_color=BG_CARD, text_color="black",
                      width=180, height=35, corner_radius=8,
                      command=self._pick_tourist_photo).pack(padx=20)
        
        right = ctk.CTkFrame(cols, fg_color=BG_WHITE, corner_radius=12)
        right.pack(side="left", expand=True, fill="both", pady=10)
        
        info = [
            f"Имя: {u['first_name']} {u['last_name']}",
            f"Никнейм: @{u['username']}",
            f"Email: {u['email']}",
            f"Статус: {u['status']}"
        ]
        for i in info:
            ctk.CTkLabel(right, text=i, font=("Arial", 14)).pack(anchor="w", padx=20, pady=6)
        
        ctk.CTkLabel(right, text="Мои бронирования:", font=("Arial", 15, "bold")).pack(anchor="w", padx=20, pady=(20, 8))
        
        bookings = DB.query(
            """SELECT e.title, b.booked_at, b.status, b.participants_count
               FROM bookings b
               JOIN schedules s ON b.schedule_id = s.id
               JOIN excursions e ON s.excursion_id = e.id
               WHERE b.tourist_id = %s
               ORDER BY b.booked_at DESC""",
            (u['id'],)
        )
        
        if bookings:
            for b in bookings:
                status_icon = {"PEND": "⏳", "CONF": "✅", "CANC": "❌", "COM": "📌"}.get(b['status'], "❓")
                ctk.CTkLabel(right, text=f"{status_icon} {b['title']} | {b['booked_at'].strftime('%d.%m.%Y') if b['booked_at'] else '—'} | {b['participants_count']} чел.", 
                            font=("Arial", 12)).pack(anchor="w", padx=35, pady=3)
        else:
            ctk.CTkLabel(right, text="Нет бронирований", text_color="#666666", font=("Arial", 12)).pack(anchor="w", padx=35, pady=8)

    def _load_tourist_photo(self, path):
        try:
            from PIL import Image as PILImage, ImageDraw
            img = PILImage.open(path).resize((200, 200)).convert("RGBA")
            mask = PILImage.new("L", (200, 200), 0)
            ImageDraw.Draw(mask).rounded_rectangle((0, 0, 200, 200), radius=20, fill=255)
            img.putalpha(mask)
            self._t_img = ctk.CTkImage(img, size=(200, 200))
            self._tourist_photo_label.configure(image=self._t_img, text="", fg_color="transparent")
        except Exception as e:
            messagebox.showerror("Ошибка фото", str(e))

    def _pick_tourist_photo(self):
        path = filedialog.askopenfilename(
            title="Выберите фото",
            filetypes=[("Изображения", "*.jpg *.jpeg *.png *.bmp *.gif")]
        )
        if not path:
            return
        import shutil
        os.makedirs("photos", exist_ok=True)
        dest = os.path.join("photos", os.path.basename(path))
        shutil.copy2(path, dest)
        DB.query("UPDATE users SET photo_path=%s WHERE id=%s",
                 (dest, self.master.current_user['id']), fetch=None)
        self._load_tourist_photo(dest)


# Управление данными
class ManageDataPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        ctk.CTkButton(self, text="← В меню", fg_color=BG_CARD, text_color="black", 
                      width=150, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)
        ctk.CTkLabel(self, text="Управление данными", font=("Arial", 22, "bold")).pack(pady=15)
        
        box_layout = ctk.CTkFrame(self, fg_color="transparent")
        box_layout.pack(fill="both", expand=True, padx=15, pady=10)
        
        # Экскурсии
        b_exc = ctk.CTkFrame(box_layout, fg_color="white", corner_radius=12)
        b_exc.pack(side="left", expand=True, fill="both", padx=(0, 10), pady=10)
        ctk.CTkLabel(b_exc, text="Экскурсии", font=("Arial", 16, "bold")).pack(pady=12)
        
        exc_list = ctk.CTkScrollableFrame(b_exc, fg_color="transparent", height=400)
        exc_list.pack(fill="both", expand=True, padx=12, pady=8)
        
        for item in DB.query("SELECT id, title, status FROM excursions ORDER BY id DESC", fetch='all'):
            row = ctk.CTkFrame(exc_list, fg_color=BG_MAIN, height=45, corner_radius=8)
            row.pack(fill="x", pady=4)
            ctk.CTkLabel(row, text=f"{item['title']} ({item['status']})", font=("Arial", 12)).pack(side="left", padx=12)
            ctk.CTkButton(row, text="🗑️", width=36, height=30, fg_color="#E57373", text_color="white", corner_radius=6,
                         command=lambda eid=item['id'], t=item['title']: self._delete_excursion(eid, t)).pack(side="right", padx=(0, 4))
            ctk.CTkButton(row, text="✏️", width=36, height=30, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda eid=item['id']: self.master.show_edit_excursion(eid)).pack(side="right", padx=4)

        ctk.CTkButton(b_exc, text="➕ Добавить экскурсию", fg_color=BG_CARD, text_color="black",
                      width=180, height=38, corner_radius=8, command=self.master.show_add_excursion).pack(side="bottom", pady=15)

        # Гиды
        b_gui = ctk.CTkFrame(box_layout, fg_color="white", corner_radius=12)
        b_gui.pack(side="left", expand=True, fill="both", padx=(10, 0), pady=10)
        ctk.CTkLabel(b_gui, text="Гиды", font=("Arial", 16, "bold")).pack(pady=12)

        guides_list = ctk.CTkScrollableFrame(b_gui, fg_color="transparent", height=400)
        guides_list.pack(fill="both", expand=True, padx=12, pady=8)

        for item in DB.query("""SELECT u.id, u.first_name, u.last_name FROM guides g JOIN users u ON u.id=g.user_id""", fetch='all'):
            row = ctk.CTkFrame(guides_list, fg_color=BG_MAIN, height=45, corner_radius=8)
            row.pack(fill="x", pady=4)
            ctk.CTkLabel(row, text=f"{item['first_name']} {item['last_name']}", font=("Arial", 12)).pack(side="left", padx=12)
            ctk.CTkButton(row, text="🗑️", width=36, height=30, fg_color="#E57373", text_color="white", corner_radius=6,
                         command=lambda gid=item['id'], n=f"{item['first_name']} {item['last_name']}": self._delete_guide(gid, n)).pack(side="right", padx=(0, 4))
            ctk.CTkButton(row, text="✏️", width=36, height=30, fg_color=BG_CARD, text_color="black", corner_radius=6,
                         command=lambda gid=item['id']: self.master.show_edit_guide(gid)).pack(side="right", padx=4)

        ctk.CTkButton(b_gui, text="➕ Создать профиль", fg_color=BG_CARD, text_color="black",
                      width=180, height=38, corner_radius=8, command=self.master.show_add_guide).pack(side="bottom", pady=15)

    def _delete_excursion(self, exc_id, title):
        if not messagebox.askyesno('Удалить', f'Удалить экскурсию "{title}"?\nВсе бронирования и отзывы также будут удалены.'):
            return
        DB.query("DELETE FROM reviews WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM bookings WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM schedules WHERE excursion_id=%s", (exc_id,), fetch=None)
        DB.query("DELETE FROM excursions WHERE id=%s", (exc_id,), fetch=None)
        messagebox.showinfo("Готово", "Экскурсия удалена.")
        self.master.show_manage_data()

    def _delete_guide(self, guide_id, name):
        if not messagebox.askyesno('Удалить', f'Удалить гида {name}?\nВсе его экскурсии, бронирования и отзывы также будут удалены.'):
            return
        excursions = DB.query("SELECT id FROM excursions WHERE guide_id=%s", (guide_id,), fetch='all') or []
        for exc in excursions:
            DB.query("DELETE FROM reviews WHERE excursion_id=%s", (exc['id'],), fetch=None)
            DB.query("DELETE FROM bookings WHERE excursion_id=%s", (exc['id'],), fetch=None)
            DB.query("DELETE FROM schedules WHERE excursion_id=%s", (exc['id'],), fetch=None)
        DB.query("DELETE FROM excursions WHERE guide_id=%s", (guide_id,), fetch=None)
        DB.query("DELETE FROM guides WHERE user_id=%s", (guide_id,), fetch=None)
        DB.query("DELETE FROM users WHERE id=%s", (guide_id,), fetch=None)
        messagebox.showinfo("Готово", f"Гид {name} удалён.")
        self.master.show_manage_data()


# Добавление/редактирование гида
class AddGuidePage(ctk.CTkFrame):
    def __init__(self, master, guide_id=None):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        self.guide_id = guide_id
        self.photo_path = None

        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black",
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)

        card = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        card.pack(fill="both", expand=True, padx=25, pady=25)

        title_txt = "Изменение гида" if guide_id else "Добавление гида"
        ctk.CTkLabel(card, text=title_txt, font=("Arial", 20, "bold")).pack(pady=20)

        content_frame = ctk.CTkFrame(card, fg_color="transparent")
        content_frame.pack(fill="both", expand=True, padx=25, pady=10)

        left = ctk.CTkFrame(content_frame, fg_color="transparent")
        left.pack(side="left", padx=25, anchor="n", pady=10)

        self.photo_label = ctk.CTkLabel(left, text="📷\nФото", width=180, height=180,
                                         fg_color=COLOR_ACCENT, corner_radius=8,
                                         font=("Arial", 14))
        self.photo_label.pack()

        ctk.CTkButton(left, text="📷 Выбрать фото", fg_color=BG_CARD, text_color="black",
                      width=180, height=35, corner_radius=8, command=self._pick_photo).pack(pady=(10, 0))

        right = ctk.CTkFrame(content_frame, fg_color="transparent")
        right.pack(side="left", expand=True, fill="both", padx=20)

        self.fields = {}
        fields = ["Имя", "Фамилия", "Email", "Специализация", "Опыт работы", "Описание"]

        for f in fields:
            row = ctk.CTkFrame(right, fg_color="transparent", height=45)
            row.pack(fill="x", pady=5)
            ctk.CTkLabel(row, text=f, width=130, anchor="w", font=("Arial", 13)).pack(side="left")
            entry = ctk.CTkEntry(row, width=320, height=38, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
            entry.pack(side="left", padx=10)
            self.fields[f] = entry

        if guide_id:
            user = DB.query("SELECT * FROM users WHERE id=%s", (guide_id,), fetch='one')
            guide = DB.query("SELECT * FROM guides WHERE user_id=%s", (guide_id,), fetch='one')
            if user:
                self.fields["Имя"].insert(0, user['first_name'] or '')
                self.fields["Фамилия"].insert(0, user['last_name'] or '')
                self.fields["Email"].insert(0, user['email'] or '')
            if guide:
                self.fields["Специализация"].insert(0, guide['specialization'] or '')
                self.fields["Опыт работы"].insert(0, str(guide['experience_years']) if guide['experience_years'] else '')
                self.fields["Описание"].insert(0, guide['description'] or '')
                if guide.get('photo_path') and os.path.exists(guide['photo_path']):
                    self.photo_path = guide['photo_path']
                    self._show_photo_preview(self.photo_path)

        ctk.CTkButton(card, text="Сохранить", fg_color=BG_CARD, text_color="black",
                      width=200, height=42, corner_radius=8, command=self._save).pack(pady=25)

    def _show_photo_preview(self, path):
        try:
            from PIL import Image as PILImage, ImageDraw
            img = PILImage.open(path).resize((180, 180)).convert("RGBA")
            mask = PILImage.new("L", (180, 180), 0)
            ImageDraw.Draw(mask).rounded_rectangle((0, 0, 180, 180), radius=20, fill=255)
            img.putalpha(mask)
            self._tk_img = ctk.CTkImage(img, size=(180, 180))
            self.photo_label.configure(image=self._tk_img, text="", fg_color="transparent")
        except Exception as e:
            messagebox.showerror("Ошибка фото", str(e))
            self.photo_label.configure(image=None, text="Ошибка")

    def _pick_photo(self):
        path = filedialog.askopenfilename(
            title="Выберите фото",
            filetypes=[("Изображения", "*.jpg *.jpeg *.png *.bmp *.gif")]
        )
        if not path:
            return
        os.makedirs("photos", exist_ok=True)
        import shutil
        filename = os.path.basename(path)
        dest = os.path.join("photos", filename)
        shutil.copy2(path, dest)
        self.photo_path = dest
        self._show_photo_preview(dest)

    def _save(self):
        name = self.fields["Имя"].get()
        surname = self.fields["Фамилия"].get()
        email = self.fields["Email"].get()
        spec = self.fields["Специализация"].get()
        exp = self.fields["Опыт работы"].get()
        desc = self.fields["Описание"].get()

        if not all([name, surname, email]):
            messagebox.showerror("Ошибка", "Заполните имя, фамилию и email")
            return

        if self.guide_id:
            DB.query(
                "UPDATE users SET first_name=%s, last_name=%s, email=%s WHERE id=%s",
                (name, surname, email, self.guide_id), fetch=None
            )
            DB.query(
                "UPDATE guides SET specialization=%s, experience_years=%s, description=%s, photo_path=%s WHERE user_id=%s",
                (spec, int(exp) if exp else 0, desc, self.photo_path, self.guide_id), fetch=None
            )
            messagebox.showinfo("Успех", "Данные гида обновлены!")
        else:
            username = email.split('@')[0]
            DB.query(
                "INSERT INTO users (username, email, password_hash, first_name, last_name, role, status)"
                " VALUES (%s,%s,%s,%s,%s,'GUI','ACT')",
                (username, email, DB.hash('guide123'), name, surname), fetch=None
            )
            user = DB.query("SELECT id FROM users WHERE email=%s", (email,), fetch='one')
            DB.query(
                "INSERT INTO guides (user_id, specialization, experience_years, description, rating, photo_path)"
                " VALUES (%s,%s,%s,%s,4.0,%s)",
                (user['id'], spec, int(exp) if exp else 0, desc, self.photo_path), fetch=None
            )
            messagebox.showinfo("Успех", f"Гид создан! Пароль: guide123")

        self.master.show_dashboard()


class EditGuidePage(AddGuidePage):
    def __init__(self, master, guide_id):
        super().__init__(master, guide_id=guide_id)


# Добавление/редактирование экскурсии
class AddExcursionPage(ctk.CTkFrame):
    def __init__(self, master, excursion_id=None):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        self.excursion_id = excursion_id
        self.photo_path = None

        ctk.CTkButton(self, text="← Назад", fg_color=BG_CARD, text_color="black",
                      width=120, height=35, corner_radius=8, command=lambda: master.show_dashboard()).pack(anchor="w", pady=10, padx=5)

        card = ctk.CTkFrame(self, fg_color="white", corner_radius=12)
        card.pack(fill="both", expand=True, padx=25, pady=25)

        title_txt = "Изменение экскурсии" if excursion_id else "Добавление экскурсии"
        ctk.CTkLabel(card, text=title_txt, font=("Arial", 20, "bold")).pack(pady=20)

        content_frame = ctk.CTkFrame(card, fg_color="transparent")
        content_frame.pack(fill="both", expand=True, padx=25, pady=10)

        left = ctk.CTkFrame(content_frame, fg_color="transparent")
        left.pack(side="left", padx=25, anchor="n", pady=10)

        self.photo_label = ctk.CTkLabel(left, text="🖼️\nФото", width=180, height=180,
                                         fg_color=COLOR_ACCENT, corner_radius=8, font=("Arial", 14))
        self.photo_label.pack()
        ctk.CTkButton(left, text="📷 Выбрать фото", fg_color=BG_CARD, text_color="black",
                      width=180, height=35, corner_radius=8, command=self._pick_photo).pack(pady=(10, 0))

        right = ctk.CTkFrame(content_frame, fg_color="transparent")
        right.pack(side="left", expand=True, fill="both", padx=20)

        self.fields = {}
        fields = ["Название тура", "Категория", "Длительность (ч)", "Макс. участников", "Описание"]
        categories = ["Историческая", "Архитектурная", "Гастрономическая", "Природная", "Обзорная"]

        for f in fields:
            row = ctk.CTkFrame(right, fg_color="transparent", height=45)
            row.pack(fill="x", pady=5)
            ctk.CTkLabel(row, text=f, width=140, anchor="w", font=("Arial", 13)).pack(side="left")
            if f == "Категория":
                entry = ctk.CTkComboBox(row, values=categories, width=300, height=38, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
            else:
                entry = ctk.CTkEntry(row, width=300, height=38, fg_color=COLOR_INPUT, corner_radius=8, border_color=BG_CARD)
            entry.pack(side="left", padx=10)
            self.fields[f] = entry

        if excursion_id:
            exc = DB.query("SELECT * FROM excursions WHERE id=%s", (excursion_id,), fetch='one')
            if exc:
                self.fields["Название тура"].insert(0, exc['title'] or '')
                self.fields["Категория"].set(exc['category'] or categories[0])
                self.fields["Длительность (ч)"].insert(0, str(exc['duration_hours']) if exc['duration_hours'] else '')
                self.fields["Макс. участников"].insert(0, str(exc['max_participants']) if exc['max_participants'] else '')
                self.fields["Описание"].insert(0, exc['description'] or '')
                if exc.get('photo_path') and os.path.exists(exc['photo_path']):
                    self.photo_path = exc['photo_path']
                    self._show_photo_preview(self.photo_path)

        ctk.CTkButton(card, text="Сохранить", fg_color=BG_CARD, text_color="black",
                      width=200, height=42, corner_radius=8, command=self._save).pack(pady=25)

    def _show_photo_preview(self, path):
        try:
            from PIL import Image as PILImage, ImageDraw
            img = PILImage.open(path).resize((180, 180)).convert("RGBA")
            mask = PILImage.new("L", (180, 180), 0)
            ImageDraw.Draw(mask).rounded_rectangle((0, 0, 180, 180), radius=20, fill=255)
            img.putalpha(mask)
            self._tk_img = ctk.CTkImage(img, size=(180, 180))
            self.photo_label.configure(image=self._tk_img, text="", fg_color="transparent")
        except Exception as e:
            messagebox.showerror("Ошибка фото", str(e))

    def _pick_photo(self):
        path = filedialog.askopenfilename(
            title="Выберите фото",
            filetypes=[("Изображения", "*.jpg *.jpeg *.png *.bmp *.gif")]
        )
        if not path:
            return
        import shutil
        os.makedirs("photos", exist_ok=True)
        dest = os.path.join("photos", os.path.basename(path))
        shutil.copy2(path, dest)
        self.photo_path = dest
        self._show_photo_preview(dest)

    def _save(self):
        title = self.fields["Название тура"].get()
        category = self.fields["Категория"].get()
        duration = self.fields["Длительность (ч)"].get()
        max_participants = self.fields["Макс. участников"].get()
        description = self.fields["Описание"].get()

        if not title:
            messagebox.showerror("Ошибка", "Введите название экскурсии")
            return

        try:
            dur = int(duration) if duration else 2
            maxp = int(max_participants) if max_participants else 10
        except:
            messagebox.showerror("Ошибка", "Длительность и макс. участников должны быть числами")
            return

        guide_id = self.master.current_user['id']

        if self.excursion_id:
            DB.query(
                "UPDATE excursions SET title=%s, description=%s, category=%s, duration_hours=%s, max_participants=%s, photo_path=%s WHERE id=%s",
                (title, description, category, dur, maxp, self.photo_path, self.excursion_id), fetch=None
            )
            messagebox.showinfo("Успех", "Экскурсия обновлена!")
        else:
            DB.query(
                "INSERT INTO excursions (guide_id, title, description, category, duration_hours, max_participants, status, photo_path)"
                " VALUES (%s,%s,%s,%s,%s,%s,'PUB',%s)",
                (guide_id, title, description, category, dur, maxp, self.photo_path), fetch=None
            )
            self.master.add_notification(
                guide_id,
                'Экскурсия добавлена',
                f'Ваша экскурсия "{title}" успешно опубликована',
                'INFO'
            )
            messagebox.showinfo("Успех", "Экскурсия добавлена!")

        self.master.show_dashboard()


class EditExcursionPage(AddExcursionPage):
    def __init__(self, master, excursion_id):
        super().__init__(master, excursion_id=excursion_id)


# Глобальный поиск (только экскурсии, без гидов)
class SearchCatalogPage(ctk.CTkFrame):
    def __init__(self, master):
        super().__init__(master, fg_color=BG_MAIN)
        self.master = master
        
        top = ctk.CTkFrame(self, fg_color="transparent")
        top.pack(fill="x", pady=8, padx=5)
        ctk.CTkButton(top, text="← В меню", fg_color=BG_CARD, text_color="black", width=120, height=35, corner_radius=8,
                      command=lambda: master.show_dashboard()).pack(side="left")
        
        self.search_entry = ctk.CTkEntry(top, placeholder_text="🔎 Поиск экскурсий...", 
                                         fg_color="white", width=450, height=40, corner_radius=8)
        self.search_entry.pack(side="left", fill="x", expand=True, padx=15)
        ctk.CTkButton(top, text="Найти", fg_color="black", text_color="white", width=100, height=40, corner_radius=8,
                      command=self._search).pack(side="left")
        
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, pady=15)
        body.grid_columnconfigure(0, weight=3)
        body.grid_columnconfigure(1, weight=1)
        body.grid_rowconfigure(0, weight=1)
        
        self.results_box = ctk.CTkFrame(body, fg_color="white", corner_radius=12)
        self.results_box.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        
        ctk.CTkLabel(self.results_box, text="Результаты поиска", font=("Arial", 16, "bold")).pack(anchor="w", padx=15, pady=12)
        self.results_container = ctk.CTkScrollableFrame(self.results_box, fg_color="transparent")
        self.results_container.pack(fill="both", expand=True, padx=12, pady=8)
        
        filter_sidebar = ctk.CTkFrame(body, fg_color=BG_CARD, corner_radius=12)
        filter_sidebar.grid(row=0, column=1, sticky="nsew")
        
        ctk.CTkLabel(filter_sidebar, text="Сортировка", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=(18, 8))
        self.sort_var = ctk.StringVar(value="По популярности")
        for opt in ["По популярности", "По рейтингу"]:
            ctk.CTkRadioButton(filter_sidebar, text=opt, variable=self.sort_var, value=opt,
                              fg_color="black", font=("Arial", 12)).pack(anchor="w", padx=25, pady=4)
        
        ctk.CTkLabel(filter_sidebar, text="Фильтры", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=(18, 8))
        
        self.filter_cat = ctk.CTkComboBox(filter_sidebar, values=["Все", "Историческая", "Архитектурная", "Гастрономическая", "Природная", "Обзорная"],
                                         fg_color="white", width=160, height=35, corner_radius=8)
        self.filter_cat.set("Все")
        self.filter_cat.pack(anchor="w", padx=20, pady=5)
        
        ctk.CTkButton(filter_sidebar, text="Применить", fg_color=BG_MAIN, text_color="black", 
                      width=140, height=38, corner_radius=8, command=self._search).pack(anchor="w", padx=20, pady=20)

    def _search(self):
        for w in self.results_container.winfo_children():
            w.destroy()
        
        query = self.search_entry.get().lower()
        category_filter = self.filter_cat.get()
        
        excursions = DB.query(
            "SELECT id, title, description, category FROM excursions WHERE status='PUB'",
            fetch='all'
        )
        
        results = []
        for e in excursions:
            if query and query not in e['title'].lower() and query not in (e['description'] or '').lower():
                continue
            if category_filter != "Все" and e['category'] != category_filter:
                continue
            results.append(("📌 Экскурсия", e['title'], e['description'] or "Нет описания"))
        
        if not results:
            ctk.CTkLabel(self.results_container, text="Ничего не найдено", text_color="gray", font=("Arial", 14)).pack(pady=30)
            return
        
        for r_type, title, desc in results:
            item = ctk.CTkFrame(self.results_container, fg_color=BG_MAIN, corner_radius=10)
            item.pack(fill="x", pady=6)
            ctk.CTkLabel(item, text=f"{r_type}: {title}", font=("Arial", 14, "bold")).pack(anchor="w", padx=15, pady=8)
            ctk.CTkLabel(item, text=desc, font=("Arial", 12), text_color="gray").pack(anchor="w", padx=15, pady=(0, 8))


# Запуск
if __name__ == "__main__":
    app = App()
    app.mainloop()